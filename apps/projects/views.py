import csv
import datetime
import json
from datetime import timedelta
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError
from django.db.models import F, Q, Sum
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import redirect, get_object_or_404, render
from django.urls import reverse
from django.views.generic import TemplateView, ListView, DetailView, View
from rest_framework import generics, status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.users.models import User
from .models import (
    Project, ProjectTemplate, PhaseTemplate, TaskTemplate,
    Phase, Task, TaskAssignment,
    ProjectMember, TaskComment, PhaseComment, ProjectMessage,
    ProjectStep, StepComment, ProjectComment,
    InstallationPlan, InstallationReport,
    ETAPES_BOISANDCO, PARTIE_LABELS,
)
from .models_pose import (
    PoseProjetMeta, Poseur, Article, Affectation, Lot, Paiement, Pointage, BonReglement,
    PoseArticleDecomposition, PosePointageComposant, PoseDemandeModification, PoseDelegation,
    PosePieceComptage, PoseDemandePaiement, PoseRecuPaiement,
)
from .permissions import IsProjectMember, IsManagerOrAbove, IsChefDeProjet
from .pose_gestion_service import (
    peut_creer_projet_pose, peut_initier_projet_pose, peut_gerer_projet_pose, peut_gerer_referentiel_poseurs,
    peut_valider_prix_pose, importer_bordereau, definir_prix_de_pose, reviser_prix_de_pose,
    creer_lot, enregistrer_pointage, lignes_reglement_semaine, verser_paiement,
    solder_poseur_semaine, solder_projet_semaine, projets_pipeline_pose, pipeline_kpis,
    definir_cumul_attachement, approuver_bon_chef, valider_bon_dg, peut_editer_bon_reglement,
    approuver_paiement, rejeter_paiement,
    BordereauImportError, PointageError,
)
from .pose_service import (
    valider_somme_poids, recap_pose, montant_article, montant_composant, avancement_article,
    releve_global_poseur, TOLERANCE_SOMME_POIDS,
)
from .pose_workflow_service import (
    projets_visibles_pointage, peut_saisir, peut_valider, peut_valider_piece, peut_gerer_decomposition,
    peut_gerer_delegation, peut_approuver_modification, peut_accorder_derogation,
    enregistrer_quantite, enregistrer_pointages_groupe, valider_pointage, valider_piece,
    refuser_piece, piece_en_attente, demander_modification, approuver_modification,
    accorder_derogation, refuser_demande, deleguer_projet, revoquer_delegation,
    article_verrouille, peut_marquer_paye, marquer_paye,
    peut_creer_demande_paiement, peut_valider_demande_paiement, peut_gerer_brouillon_demande,
    obtenir_ou_creer_brouillon, ajouter_lignes_demande, retirer_ligne_demande,
    soumettre_demande_paiement, fiches_engagees_ids,
    valider_demande_paiement, refuser_demande_paiement,
    peut_gerer_tresorerie_pose, lignes_validees_non_payees, creer_recu_paiement,
    PointageWorkflowError,
)
from .serializers_comments import (
    StepCommentSerializer, ProjectCommentSerializer, ProjectStepSerializer,
)
from .services import desig_chef_de_projet, validate_chef, complete_step


def _visible_projects(user):
    """Projects the user is allowed to see.
    DIRECTEUR, MANAGER, ESTIMATEUR, ADV, ATELIER and superusers see everything.
    ADMIN and DG_POSE additionally see every projet de pose (module Gestion de la pose),
    matching their DG-level oversight there (peut_gerer_projet_pose).
    Others see projects where they have any of the following roles:
    - manager or estimator on the project
    - declared ProjectMember
    - designated chef de projet
    - created the project directly (created_by)
    - commercial who submitted the chiffrage that originated the project
    """
    if user.is_superuser or user.role in ('DIRECTEUR', 'MANAGER', 'ESTIMATEUR', 'ADV', 'ATELIER'):
        return Project.objects.all()
    qs = Project.objects.filter(
        Q(manager=user) |
        Q(estimator=user) |
        Q(members__user=user) |
        Q(chef_de_projet=user) |
        Q(created_by=user) |
        Q(demande_chiffrage__commercial=user)
    )
    if user.role in ('ADMIN', 'DG_POSE'):
        qs = qs | Project.objects.filter(pose_meta__isnull=False)
    return qs.distinct()


# Maps phase.owner_label → list of User.role codes who can comment
_OWNER_LABEL_ROLES = {
    'Commercial':       ['COMMERCIAL'],
    'DC + Commercial':  ['MANAGER', 'COMMERCIAL'],
    'DC':               ['MANAGER'],
    'BE':               ['ESTIMATEUR'],
    'BE + DG':          ['ESTIMATEUR', 'DIRECTEUR'],
    'BE Méthodes':      ['ESTIMATEUR'],
    'ADV':              ['ADV'],
    'Production':       ['ATELIER'],
    'Logistique':       ['CHAUFFEUR'],
    'Equipe pose':      ['POSEUR'],
    'Pose':             ['POSEUR'],
    'Pose + ADV':       ['POSEUR', 'ADV'],
    'ADV + Finance':    ['ADV', 'FINANCE'],
    'ADV + Pose':       ['ADV', 'POSEUR'],
    'SAV':              ['SAV'],
}


def _can_comment_phase(user, phase):
    """True if user's role matches the phase owner or user is management."""
    if user.is_superuser or user.role in ('DIRECTEUR', 'MANAGER'):
        return True
    allowed = _OWNER_LABEL_ROLES.get(phase.owner_label, [])
    return user.role in allowed


class ProjectAccessMixin(LoginRequiredMixin):
    """Mixin that raises 404 for users without access to the project."""

    def _get_project(self):
        pk = self.kwargs.get('pk') or self.kwargs.get('project_pk')
        project = get_object_or_404(Project, pk=pk)
        if not _visible_projects(self.request.user).filter(pk=project.pk).exists():
            raise Http404
        return project

    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_authenticated:
            from django.conf import settings
            from django.contrib.auth.views import redirect_to_login
            return redirect_to_login(request.get_full_path())
        # Validate project access early
        try:
            self._get_project()
        except Http404:
            raise
        return super().dispatch(request, *args, **kwargs)


class DashboardView(LoginRequiredMixin, TemplateView):
    """Vue tableau de bord — point d'entrée de l'application."""
    template_name = 'projects/dashboard.html'

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated and getattr(request.user, 'role', None) in (
            'DG_POSE', 'CHEF_PROJET_POSE', 'TRESO_POSE',
        ):
            return redirect('projects:pose_gestion_pipeline')
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        from django.db.models import Sum, Count, Q
        from apps.chiffrage.models import DemandeChiffrage

        ctx = super().get_context_data(**kwargs)
        user = self.request.user
        all_projects = _visible_projects(user).select_related('manager')
        active = all_projects.exclude(status=Project.Status.CLOTURE)

        ctx['kpis'] = {
            'total': all_projects.count(),
            'actifs': active.count(),
            'en_production': all_projects.filter(status=Project.Status.PRODUCTION).count(),
            'en_pose': all_projects.filter(status=Project.Status.POSE).count(),
        }
        ctx['projets_actifs'] = active.order_by('-created_at')[:8]

        pipeline = [
            {'status': s, 'label': l, 'count': all_projects.filter(status=s).count()}
            for s, l in Project.Status.choices
        ]
        ctx['pipeline'] = pipeline

        ctx['chart_projets'] = json.dumps({
            'labels': [p['label'] for p in pipeline],
            'data':   [p['count'] for p in pipeline],
            'colors': ['#adb5bf', '#3b82f6', '#f59e0b', '#a855f7', '#22804c', '#d1d5db'],
        })

        # ── KPI Chiffrage ────────────────────────────────────────────
        ROLES_CHIFFRAGE = ('COMMERCIAL', 'MANAGER', 'DIRECTEUR', 'ESTIMATEUR', 'RESP_METHODES')
        if user.is_superuser or user.role in ROLES_CHIFFRAGE:
            all_chiffrage = DemandeChiffrage.objects.all()

            if user.role == 'COMMERCIAL' and not user.is_superuser:
                chiffrage_qs = all_chiffrage.filter(commercial=user)
            elif user.role == 'ESTIMATEUR' and not user.is_superuser:
                chiffrage_qs = all_chiffrage.filter(assigned_to=user)
            else:
                chiffrage_qs = all_chiffrage

            POST_DG = ['DEVIS_VALIDE', 'TRANSMIS', 'ACCEPTE']
            montant_total = chiffrage_qs.filter(statut__in=POST_DG).aggregate(
                t=Sum('montant_ht')
            )['t'] or 0

            closes  = chiffrage_qs.filter(statut__in=['ACCEPTE', 'REFUSE_CLI']).count()
            acceptes = chiffrage_qs.filter(statut='ACCEPTE').count()

            from django.utils import timezone as tz
            today      = tz.now().date()
            CLOSED_ST  = ['DEVIS_VALIDE', 'TRANSMIS', 'ACCEPTE', 'REFUSE_CLI', 'ARCHIVE']

            n_attente  = chiffrage_qs.filter(statut__in=['EN_ATTENTE', 'RETOURNEE', 'REJETEE']).count()
            n_cours    = chiffrage_qs.filter(statut__in=['VALIDEE_DC', 'EN_CHIFFRAGE', 'SOUMIS_RM', 'SOUMIS_DG', 'REVISION_DG']).count()
            n_valides  = chiffrage_qs.filter(statut__in=POST_DG).count()
            n_refuses  = chiffrage_qs.filter(statut__in=['REFUSE_CLI', 'ARCHIVE']).count()
            n_retards  = chiffrage_qs.filter(
                delai_souhaite__isnull=False,
                delai_souhaite__lt=today,
            ).exclude(statut__in=CLOSED_ST).count()

            m = float(montant_total)
            if m >= 1_000_000:
                montant_fmt = f"{m / 1_000_000:.2f} M MAD"
            elif m >= 1_000:
                montant_fmt = f"{m / 1_000:.0f} K MAD"
            elif m > 0:
                montant_fmt = f"{m:,.0f} MAD"
            else:
                montant_fmt = None

            ctx['kpi_chiffrage'] = {
                'en_attente':  n_attente,
                'en_cours':    n_cours,
                'valides':     n_valides,
                'montant_fmt': montant_fmt,
                'retards':     n_retards,
                'taux':        round(acceptes / closes * 100) if closes else None,
                'acceptes':    acceptes,
                # Données graphiques
                'chart_donut': json.dumps({
                    'labels': ['En attente', 'En cours', 'Devis validés', 'Refusés / Archivés'],
                    'data':   [n_attente, n_cours, n_valides, n_refuses],
                    'colors': ['#f59e0b', '#6366f1', '#22804c', '#adb5bf'],
                }),
            }

            # Performance par méthodiste (managers, RM, DG, admin)
            if user.is_superuser or user.role in ('DIRECTEUR', 'MANAGER', 'RESP_METHODES'):
                methodes_perf = list(
                    all_chiffrage.filter(assigned_to__isnull=False)
                    .values('assigned_to__pk', 'assigned_to__first_name', 'assigned_to__last_name')
                    .annotate(
                        total=Count('pk'),
                        en_cours=Count('pk', filter=Q(statut__in=['EN_CHIFFRAGE', 'SOUMIS_RM', 'REVISION_DG'])),
                        valides=Count('pk', filter=Q(statut__in=['SOUMIS_DG', 'DEVIS_VALIDE', 'TRANSMIS', 'ACCEPTE'])),
                        montant=Sum('montant_ht', filter=Q(montant_ht__isnull=False)),
                    )
                    .order_by('-valides', '-total')
                )
                ctx['methodes_perf'] = methodes_perf
                ctx['kpi_chiffrage']['chart_methodes'] = json.dumps({
                    'labels':  [f"{m['assigned_to__first_name']} {m['assigned_to__last_name']}" for m in methodes_perf],
                    'en_cours': [m['en_cours'] for m in methodes_perf],
                    'valides':  [m['valides'] for m in methodes_perf],
                    'montant':  [float(m['montant'] or 0) for m in methodes_perf],
                })

        return ctx


class ProjectListView(LoginRequiredMixin, ListView):
    model = Project
    context_object_name = 'projects'

    def dispatch(self, request, *args, **kwargs):
        from apps.users.models import ROLES_POSE_ONLY
        if getattr(request.user, 'role', None) in ROLES_POSE_ONLY:
            return redirect('projects:pose_gestion_pipeline')
        return super().dispatch(request, *args, **kwargs)

    def get_template_names(self):
        if self.request.GET.get('partial') == 'list':
            return ['projects/_project_list_partial.html']
        return ['projects/project_list.html']

    def get_queryset(self):
        qs = _visible_projects(self.request.user).select_related('manager')

        # Recherche textuelle : nom, référence, client, adresse
        q = self.request.GET.get('q', '').strip()
        if q:
            qs = qs.filter(
                Q(name__icontains=q) |
                Q(reference__icontains=q) |
                Q(client_name__icontains=q) |
                Q(address__icontains=q)
            )

        # Filtre statut
        status = self.request.GET.get('status', '')
        if status:
            qs = qs.filter(status=status)

        # Filtre manager
        manager_id = self.request.GET.get('manager', '')
        if manager_id:
            try:
                qs = qs.filter(manager_id=int(manager_id))
            except (ValueError, TypeError):
                pass

        # Filtre date début ≥
        start_after = self.request.GET.get('start_after', '')
        if start_after:
            qs = qs.filter(start_date__gte=start_after)

        # Filtre date fin ≤
        end_before = self.request.GET.get('end_before', '')
        if end_before:
            qs = qs.filter(end_date__lte=end_before)

        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['project_statuses'] = Project.Status.choices
        ctx['managers'] = (
            User.objects.filter(is_active_employee=True)
            .order_by('last_name', 'first_name')
        )
        # Valeurs courantes des filtres
        ctx['q']               = self.request.GET.get('q', '')
        ctx['selected_status'] = self.request.GET.get('status', '')
        ctx['selected_manager']= self.request.GET.get('manager', '')
        ctx['start_after']     = self.request.GET.get('start_after', '')
        ctx['end_before']      = self.request.GET.get('end_before', '')
        # Nombre de filtres actifs
        ctx['active_filters_count'] = sum([
            bool(ctx['q']),
            bool(ctx['selected_status']),
            bool(ctx['selected_manager']),
            bool(ctx['start_after']),
            bool(ctx['end_before']),
        ])
        # Total non filtré (pour affichage "X / Y projets")
        ctx['total_count'] = _visible_projects(self.request.user).count()
        return ctx


class ProjectCreateView(LoginRequiredMixin, View):
    """Création manuelle de projet — réservée aux administrateurs.
    Tous les autres rôles créent des projets via le workflow chiffrage.
    """
    template_name = 'projects/project_form.html'

    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_superuser:
            from django.core.exceptions import PermissionDenied
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def _ctx(self):
        team = User.objects.filter(is_active_employee=True).order_by('last_name')
        by_role = {}
        for u in team:
            by_role.setdefault(u.get_role_display(), []).append(u)
        return {
            'teams': team,
            'teams_by_role': by_role,
            'templates': ProjectTemplate.objects.filter(is_active=True),
            'statuses': Project.Status.choices,
            'errors': {},
            'post': {},
        }

    def get(self, request):
        return render(request, self.template_name, self._ctx())

    def post(self, request):
        data = request.POST
        files = request.FILES
        errors = {}

        name = data.get('name', '').strip()
        client_name = data.get('client_name', '').strip()
        manager_id = data.get('manager')
        status = data.get('status', Project.Status.AVANT_VENTE)

        if not name:
            errors['name'] = 'Le nom du projet est obligatoire.'
        if not client_name:
            errors['client_name'] = 'Le nom du client est obligatoire.'
        if not manager_id:
            errors['manager'] = 'Veuillez désigner un manager.'

        if errors:
            ctx = self._ctx()
            ctx.update({'errors': errors, 'post': data})
            return render(request, self.template_name, ctx)

        project = Project.objects.create(
            name=name,
            client_name=client_name,
            client_email=data.get('client_email', ''),
            client_phone=data.get('client_phone', ''),
            address=data.get('address', ''),
            manager_id=manager_id,
            estimator_id=data.get('estimator') or None,
            status=status,
            start_date=data.get('start_date') or None,
            end_date=data.get('end_date') or None,
            budget=data.get('budget') or None,
            notes=data.get('notes', ''),
            created_by=request.user,
            initial_order=files.get('initial_order') or None,
        )

        # Membres de l'équipe sélectionnés (checkboxes)
        member_ids = data.getlist('members')
        for uid in member_ids:
            ProjectMember.objects.get_or_create(
                project=project,
                user_id=uid,
                defaults={'added_by': request.user},
            )

        # Instancier depuis gabarit si sélectionné
        template_id = data.get('template')
        if template_id:
            try:
                tmpl = ProjectTemplate.objects.get(pk=template_id)
                project.create_from_template(tmpl)
            except ProjectTemplate.DoesNotExist:
                pass
        else:
            # Phases et tâches saisies manuellement (JSON)
            phases_json = data.get('phases_json', '').strip()
            if phases_json:
                try:
                    phases_data = json.loads(phases_json)
                    for i, ph in enumerate(phases_data):
                        phase_name = ph.get('name', '').strip()
                        if not phase_name:
                            continue
                        phase = Phase.objects.create(
                            project=project,
                            name=phase_name,
                            order=i,
                            estimated_days=int(ph.get('days') or 0),
                            is_active=(i == 0),
                        )
                        for j, tk in enumerate(ph.get('tasks', [])):
                            task_name = tk.get('name', '').strip()
                            if not task_name:
                                continue
                            Task.objects.create(
                                phase=phase,
                                name=task_name,
                                order=j,
                                required_role=tk.get('role', 'ANY'),
                                estimated_hours=float(tk.get('hours') or 0),
                                status=(
                                    Task.Status.ACTIVE
                                    if i == 0 and j == 0
                                    else Task.Status.EN_ATTENTE
                                ),
                            )
                except (json.JSONDecodeError, ValueError, KeyError):
                    pass

        messages.success(request, f'Projet « {project.name} » créé avec succès.')
        return redirect('projects:detail', pk=project.pk)


class TaskDetailView(LoginRequiredMixin, DetailView):
    """Détail d'une tâche : assignés, commentaires, pièces jointes."""
    model = Task
    template_name = 'projects/task_detail.html'
    context_object_name = 'task'

    def get_queryset(self):
        return Task.objects.select_related(
            'phase__project__manager',
        ).prefetch_related(
            'assignments__user',
            'comments__author',
        )

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        task = self.object
        ctx['project'] = task.phase.project
        ctx['phase'] = task.phase
        ctx['assignable_users'] = User.objects.filter(is_active_employee=True).order_by('last_name')
        ctx['assigned_ids'] = list(task.assignments.values_list('user_id', flat=True))
        return ctx


class TaskCommentCreateView(LoginRequiredMixin, View):
    """Crée un commentaire (texte + fichier + audio) sur une tâche."""

    def post(self, request, pk):
        task = get_object_or_404(Task, pk=pk)
        text = request.POST.get('text', '').strip()
        attachment = request.FILES.get('attachment')
        audio = request.FILES.get('audio')

        if text or attachment or audio:
            TaskComment.objects.create(
                task=task,
                author=request.user,
                text=text,
                attachment=attachment or None,
                audio=audio or None,
            )
        return redirect('projects:task_detail', pk=task.pk)


class PhaseCommentCreateView(LoginRequiredMixin, View):
    """Crée un commentaire (texte + fichier + audio) sur une phase du planning.
    Seuls les rôles responsables de la phase (et les managers) peuvent commenter.
    """

    def post(self, request, pk):
        phase = get_object_or_404(Phase, pk=pk)
        if not _can_comment_phase(request.user, phase):
            from django.http import HttpResponseForbidden
            return HttpResponseForbidden('Vous ne pouvez pas commenter cette étape.')
        text = request.POST.get('text', '').strip()
        attachment = request.FILES.get('attachment')
        audio = request.FILES.get('comment_audio')

        if text or attachment or audio:
            PhaseComment.objects.create(
                phase=phase,
                author=request.user,
                text=text,
                attachment=attachment or None,
                audio=audio or None,
            )
        return redirect('projects:detail', pk=phase.project_id)


class ProjectMessageCreateView(LoginRequiredMixin, View):
    """Crée un message dans le chat projet (HTMX)."""

    def post(self, request, pk):
        project = get_object_or_404(Project, pk=pk)
        content = request.POST.get('content', '').strip()
        attachment = request.FILES.get('attachment')
        audio = request.FILES.get('audio')
        reply_to_id = request.POST.get('reply_to') or None

        reply_to = None
        if reply_to_id:
            try:
                reply_to = ProjectMessage.objects.get(pk=reply_to_id, project=project)
            except ProjectMessage.DoesNotExist:
                pass

        if content or attachment or audio:
            ProjectMessage.objects.create(
                project=project,
                author=request.user,
                content=content,
                reply_to=reply_to,
                attachment=attachment or None,
                attachment_name=attachment.name if attachment else '',
                audio=audio or None,
            )

        msgs = (
            ProjectMessage.objects
            .filter(project=project)
            .select_related('author', 'reply_to__author')
            .order_by('created_at')
        )
        return render(request, 'projects/partials/_chat_messages.html', {
            'chat_messages': msgs,
            'project': project,
            'request': request,
        })


class ProjectMessageListView(LoginRequiredMixin, View):
    """Rafraîchit les messages du chat (polling HTMX)."""

    def get(self, request, pk):
        project = get_object_or_404(Project, pk=pk)
        msgs = (
            ProjectMessage.objects
            .filter(project=project)
            .select_related('author', 'reply_to__author')
            .order_by('created_at')
        )
        return render(request, 'projects/partials/_chat_messages.html', {
            'chat_messages': msgs,
            'project': project,
            'request': request,
        })


class PhaseActivateView(LoginRequiredMixin, View):
    """Active une phase (via HTMX POST)."""

    def post(self, request, pk):
        phase = get_object_or_404(Phase, pk=pk)
        if not phase.is_active and not phase.is_completed:
            phase.activate()
        return redirect('projects:detail', pk=phase.project_id)


class PhaseCompleteView(LoginRequiredMixin, View):
    """Clôture une phase et active la suivante (via HTMX POST)."""

    def post(self, request, pk):
        phase = get_object_or_404(Phase, pk=pk)
        if phase.is_active and not phase.is_completed:
            phase.complete()
        return redirect('projects:detail', pk=phase.project_id)


class TaskProgressView(LoginRequiredMixin, View):
    """Met à jour l'avancement d'une tâche + crée un commentaire optionnel."""

    def post(self, request, pk):
        task = get_object_or_404(Task, pk=pk)
        try:
            value = int(request.POST.get('progress', 0))
        except (ValueError, TypeError):
            value = 0
        task.set_progress(value)

        # Commentaire optionnel joint à la mise à jour
        text = request.POST.get('comment_text', '').strip()
        attachment = request.FILES.get('comment_file')
        audio = request.FILES.get('comment_audio')
        if text or attachment or audio:
            TaskComment.objects.create(
                task=task,
                author=request.user,
                text=text,
                attachment=attachment or None,
                audio=audio or None,
            )

        return redirect('projects:detail', pk=task.phase.project_id)


class TaskAssignView(LoginRequiredMixin, View):
    """Affecte un utilisateur à une tâche (via HTMX POST)."""

    def post(self, request, pk):
        task = get_object_or_404(Task, pk=pk)
        user_id = request.POST.get('user_id')
        if user_id:
            TaskAssignment.objects.get_or_create(
                task=task,
                user_id=user_id,
                defaults={'assigned_by': request.user, 'is_primary': True},
            )
        return redirect('projects:detail', pk=task.phase.project_id)


class ProjectDetailView(LoginRequiredMixin, DetailView):
    model = Project
    template_name = 'projects/project_detail.html'
    context_object_name = 'project'

    def get_queryset(self):
        return _visible_projects(self.request.user).prefetch_related(
            'phases__tasks__assignments__user',
            'phases__tasks__comments',
            'phases__comments__author',
            'members__user',
        ).select_related('manager', 'estimator', 'chef_de_projet', 'chef_designe_par')

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        project = self.object
        user = self.request.user
        phases = list(project.phases.order_by('order'))

        parties = {
            'A': {'label': PARTIE_LABELS['A'], 'color': 'foret',  'phases': []},
            'B': {'label': PARTIE_LABELS['B'], 'color': 'blue',   'phases': []},
            'C': {'label': PARTIE_LABELS['C'], 'color': 'amber',  'phases': []},
        }
        for ph in phases:
            key = ph.partie if ph.partie in parties else 'A'
            parties[key]['phases'].append(ph)

        # Attach comment permission directly to each phase object
        for ph in phases:
            ph.user_can_comment = _can_comment_phase(user, ph)

        ctx['parties']    = parties
        ctx['step_range'] = range(1, 19)
        ctx['chat_messages'] = (
            ProjectMessage.objects
            .filter(project=project)
            .select_related('author', 'reply_to__author')
            .order_by('created_at')
        )

        # Chef de projet
        ctx['can_designate_chef'] = user.is_superuser or user.role in ('DIRECTEUR', 'MANAGER')
        ctx['designatable_users'] = (
            User.objects.filter(is_active_employee=True).order_by('last_name', 'first_name')
            if ctx['can_designate_chef'] else []
        )

        # Project steps with per-step comment permission
        steps = list(
            ProjectStep.objects
            .filter(project=project)
            .prefetch_related('comments__author')
            .order_by('order')
        )
        for s in steps:
            s.user_can_comment = (
                user.is_superuser
                or user.role in ('DIRECTEUR', 'MANAGER')
                or user.role in (s.responsables_roles or [])
            )
        ctx['project_steps'] = steps
        return ctx


# ---------------------------------------------------------------------------
# Gestion des gabarits (Templates)
# ---------------------------------------------------------------------------

class TemplateListView(LoginRequiredMixin, ListView):
    """Liste tous les gabarits de projet."""
    model = ProjectTemplate
    template_name = 'projects/template_list.html'
    context_object_name = 'templates'
    ordering = ['name']

    def get_queryset(self):
        return ProjectTemplate.objects.prefetch_related(
            'phase_templates__task_templates'
        ).order_by('name')


class TemplateSaveView(LoginRequiredMixin, View):
    """Créer ou modifier un gabarit via formulaire Alpine.js dynamique."""
    template_name = 'projects/template_form.html'

    def _build_phases_json(self, template):
        phases = []
        for pt in template.phase_templates.order_by('order'):
            tasks = [
                {
                    'name': tt.name,
                    'role': tt.required_role,
                    'hours': str(tt.estimated_hours),
                }
                for tt in pt.task_templates.order_by('order')
            ]
            phases.append({
                'name': pt.name,
                'days': pt.estimated_days,
                'tasks': tasks,
            })
        return json.dumps(phases)

    def get(self, request, pk=None):
        tmpl = get_object_or_404(ProjectTemplate, pk=pk) if pk else None
        ctx = {
            'tmpl': tmpl,
            'phases_json': self._build_phases_json(tmpl) if tmpl else '[]',
        }
        return render(request, self.template_name, ctx)

    def post(self, request, pk=None):
        name = request.POST.get('name', '').strip()
        description = request.POST.get('description', '').strip()
        is_active = request.POST.get('is_active') == 'on'

        if not name:
            tmpl = get_object_or_404(ProjectTemplate, pk=pk) if pk else None
            ctx = {
                'tmpl': tmpl,
                'phases_json': self._build_phases_json(tmpl) if tmpl else '[]',
                'error': 'Le nom du gabarit est obligatoire.',
            }
            return render(request, self.template_name, ctx)

        if pk:
            tmpl = get_object_or_404(ProjectTemplate, pk=pk)
            tmpl.name = name
            tmpl.description = description
            tmpl.is_active = is_active
            tmpl.save()
            # Supprimer les anciennes phases (cascade supprime les tâches)
            tmpl.phase_templates.all().delete()
            action = 'mis à jour'
        else:
            tmpl = ProjectTemplate.objects.create(
                name=name,
                description=description,
                is_active=is_active,
                created_by=request.user,
            )
            action = 'créé'

        # Recréer phases et tâches depuis le JSON Alpine.js
        phases_json = request.POST.get('phases_json', '[]').strip()
        try:
            phases_data = json.loads(phases_json)
            for i, ph in enumerate(phases_data):
                phase_name = ph.get('name', '').strip()
                if not phase_name:
                    continue
                phase = PhaseTemplate.objects.create(
                    template=tmpl,
                    name=phase_name,
                    order=i,
                    estimated_days=int(ph.get('days') or 0),
                )
                for j, tk in enumerate(ph.get('tasks', [])):
                    task_name = tk.get('name', '').strip()
                    if not task_name:
                        continue
                    TaskTemplate.objects.create(
                        phase=phase,
                        name=task_name,
                        order=j,
                        required_role=tk.get('role', 'ANY'),
                        estimated_hours=float(tk.get('hours') or 0),
                    )
        except (json.JSONDecodeError, ValueError, KeyError):
            pass

        messages.success(request, f'Gabarit « {tmpl.name} » {action} avec succès.')
        return redirect('projects:template_list')


class TemplateDeleteView(LoginRequiredMixin, View):
    """Supprime un gabarit."""

    def post(self, request, pk):
        tmpl = get_object_or_404(ProjectTemplate, pk=pk)
        name = tmpl.name
        tmpl.delete()
        messages.success(request, f'Gabarit « {name} » supprimé.')
        return redirect('projects:template_list')


# ---------------------------------------------------------------------------
# Gestion de la pose
# ---------------------------------------------------------------------------

class PoseDashboardView(LoginRequiredMixin, TemplateView):
    """Dashboard pose : tous les projets en phase POSE avec leur équipe."""
    template_name = 'projects/pose_dashboard.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        pose_projects = Project.objects.filter(
            status=Project.Status.POSE
        ).select_related('manager').prefetch_related(
            'installation_plan__chef_chantier',
            'installation_plan__team_members',
            'installation_reports',
        )
        # Tous les projets actifs pour vue globale
        all_active = Project.objects.exclude(
            status=Project.Status.CLOTURE
        ).select_related('manager').prefetch_related('installation_plan')

        poseurs = User.objects.filter(
            role__in=[User.Role.POSEUR, User.Role.CHAUFFEUR, User.Role.MANAGER],
            is_active_employee=True,
        ).order_by('last_name')

        ctx['pose_projects'] = pose_projects
        ctx['all_active']    = all_active
        ctx['poseurs']       = poseurs
        ctx['total_en_pose'] = pose_projects.count()
        return ctx


class InstallationPlanUpdateView(LoginRequiredMixin, View):
    """Crée ou met à jour le plan de pose d'un projet."""

    def post(self, request, pk):
        project = get_object_or_404(Project, pk=pk)
        plan, _ = InstallationPlan.objects.get_or_create(project=project)

        plan.chef_chantier_id = request.POST.get('chef_chantier') or None
        plan.planned_start    = request.POST.get('planned_start') or None
        plan.planned_end      = request.POST.get('planned_end') or None
        plan.notes            = request.POST.get('notes', '').strip()
        plan.save()

        member_ids = request.POST.getlist('team_members')
        plan.team_members.set(member_ids)

        messages.success(request, 'Plan de pose mis à jour.')
        return redirect('projects:pose_project', pk=project.pk)


class PoseProjectView(LoginRequiredMixin, DetailView):
    """Détail pose d'un projet : équipe + rapports."""
    model = Project
    template_name = 'projects/pose_project.html'
    context_object_name = 'project'

    def get_queryset(self):
        return Project.objects.select_related('manager').prefetch_related(
            'installation_plan__chef_chantier',
            'installation_plan__team_members',
            'installation_reports__reported_by',
        )

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        project = self.object
        try:
            ctx['plan'] = project.installation_plan
        except InstallationPlan.DoesNotExist:
            ctx['plan'] = None

        ctx['reports'] = project.installation_reports.select_related('reported_by').order_by('-year', '-week_number')
        ctx['poseurs']  = User.objects.filter(
            role__in=[User.Role.POSEUR, User.Role.CHAUFFEUR, User.Role.MANAGER],
            is_active_employee=True,
        ).order_by('last_name')

        import datetime
        today = datetime.date.today()
        ctx['current_week'] = today.isocalendar()[1]
        ctx['current_year'] = today.year
        return ctx


class InstallationReportCreateView(LoginRequiredMixin, View):
    """Crée un rapport hebdomadaire de pose."""

    def post(self, request, pk):
        project = get_object_or_404(Project, pk=pk)
        import datetime
        today = datetime.date.today()

        week = int(request.POST.get('week_number', today.isocalendar()[1]))
        year = int(request.POST.get('year', today.year))

        report, created = InstallationReport.objects.get_or_create(
            project=project,
            week_number=week,
            year=year,
            defaults={'reported_by': request.user},
        )
        if not created:
            report.reported_by = request.user

        report.progress_percent      = int(request.POST.get('progress_percent', 0))
        report.notes                 = request.POST.get('notes', '').strip()
        report.difficulties          = request.POST.get('difficulties', '').strip()
        report.incident              = bool(request.POST.get('incident'))
        report.incident_description  = request.POST.get('incident_description', '').strip()

        photo = request.FILES.get('photo')
        if photo:
            report.photo = photo

        report.save()
        messages.success(request, f'Rapport S{week}/{year} enregistré.')
        return redirect('projects:pose_project', pk=project.pk)


# ---------------------------------------------------------------------------
# Gestion de la pose — extension (CDC Nisrine ALAMI v1.0, 20/07/2026)
# Étape 1 : création projet de pose + import bordereau .xlsx
# ---------------------------------------------------------------------------

class PoseGestionCreateProjetView(LoginRequiredMixin, View):
    """Étape 1 du workflow pose : création du projet + import du bordereau .xlsx."""

    template_name = 'projects/pose_gestion_create.html'

    def dispatch(self, request, *args, **kwargs):
        if not peut_initier_projet_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def _ctx(self):
        team = User.objects.filter(is_active_employee=True).order_by('last_name')
        return {
            'chefs': team,
            'commerciaux': team.filter(role=User.Role.COMMERCIAL),
            'coordinateurs': team.filter(role='CONDUCTEUR_POSE'),
            'errors': {},
            'post': {},
            'selected_coordinateurs': [],
        }

    def get(self, request):
        return render(request, self.template_name, self._ctx())

    def post(self, request):
        data = request.POST
        fichier = request.FILES.get('bordereau')
        coordinateur_ids = request.POST.getlist('coordinateurs')
        errors = {}

        name = data.get('name', '').strip()
        client_name = data.get('client_name', '').strip()
        region = data.get('region', '').strip()
        chef_de_projet_id = data.get('chef_de_projet')
        commercial_id = data.get('commercial')
        start_date = data.get('start_date') or None

        if not name:
            errors['name'] = 'Le nom du projet est obligatoire.'
        if not client_name:
            errors['client_name'] = 'Le client est obligatoire.'
        if not region:
            errors['region'] = 'La région est obligatoire.'
        if not chef_de_projet_id:
            errors['chef_de_projet'] = 'Le chef de projet est obligatoire.'
        if not commercial_id:
            errors['commercial'] = 'Le commercial est obligatoire.'
        if not start_date:
            errors['start_date'] = 'La date de début est obligatoire.'
        if not fichier:
            errors['bordereau'] = 'Le bordereau (.xlsx) est obligatoire.'

        if errors:
            ctx = self._ctx()
            ctx.update({'errors': errors, 'post': data, 'selected_coordinateurs': coordinateur_ids})
            return render(request, self.template_name, ctx)

        project = Project.objects.create(
            name=name,
            client_name=client_name,
            manager_id=commercial_id,
            chef_de_projet_id=chef_de_projet_id,
            start_date=start_date,
            status=Project.Status.POSE,
            created_by=request.user,
        )
        PoseProjetMeta.objects.create(
            project=project,
            region=region,
            commercial_id=commercial_id,
        )

        try:
            nb_crees, nb_ignorees = importer_bordereau(project, fichier)
        except BordereauImportError as exc:
            project.delete()
            ctx = self._ctx()
            ctx.update({'errors': {'bordereau': str(exc)}, 'post': data, 'selected_coordinateurs': coordinateur_ids})
            return render(request, self.template_name, ctx)

        # project.chef_de_projet_id garde le type str de la valeur POST tant que l'instance
        # n'est pas rechargée — peut_gerer_delegation le compare à user.id (int) : sans ce
        # refresh, la comparaison échoue toujours et deleguer_projet refuse à tort.
        project.refresh_from_db()

        nb_coordinateurs = 0
        for coordinateur in User.objects.filter(pk__in=coordinateur_ids, role='CONDUCTEUR_POSE'):
            try:
                deleguer_projet(request.user, project, coordinateur)
                nb_coordinateurs += 1
            except PointageWorkflowError:
                pass

        msg = f'Projet de pose « {project.name} » créé — {nb_crees} article(s) importé(s).'
        if nb_ignorees:
            msg += f' ({nb_ignorees} ligne(s) ignorée(s), ex. totaux)'
        if nb_coordinateurs:
            msg += f' — {nb_coordinateurs} coordinateur(s) de la pose affecté(s).'
        messages.success(request, msg)
        return redirect('projects:pose_gestion_projet', pk=project.pk)


class PoseGestionProjetUpdateView(LoginRequiredMixin, View):
    """Modification des informations d'un projet de pose (accessible depuis le Pipeline)."""

    template_name = 'projects/pose_gestion_projet_edit.html'

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def _ctx(self):
        team = User.objects.filter(is_active_employee=True).order_by('last_name')
        meta = getattr(self.project, 'pose_meta', None)
        return {
            'project': self.project,
            'chefs': team,
            'commerciaux': team.filter(role=User.Role.COMMERCIAL),
            'coordinateurs': team.filter(role='CONDUCTEUR_POSE'),
            'errors': {},
            'post': {
                'name': self.project.name,
                'client_name': self.project.client_name,
                'region': meta.region if meta else '',
                'chef_de_projet': str(self.project.chef_de_projet_id or ''),
                'commercial': str(meta.commercial_id or '') if meta else '',
                'start_date': self.project.start_date.isoformat() if self.project.start_date else '',
            },
            'selected_coordinateurs': [
                str(pk) for pk in self.project.pose_delegations.filter(actif=True)
                .values_list('conducteur_delegue_id', flat=True)
            ],
        }

    def get(self, request, pk):
        return render(request, self.template_name, self._ctx())

    def post(self, request, pk):
        data = request.POST
        errors = {}

        name = data.get('name', '').strip()
        client_name = data.get('client_name', '').strip()
        region = data.get('region', '').strip()
        chef_de_projet_id = data.get('chef_de_projet')
        commercial_id = data.get('commercial')
        start_date = data.get('start_date') or None

        if not name:
            errors['name'] = 'Le nom du projet est obligatoire.'
        if not client_name:
            errors['client_name'] = 'Le client est obligatoire.'
        if not region:
            errors['region'] = 'La région est obligatoire.'
        if not chef_de_projet_id:
            errors['chef_de_projet'] = 'Le chef de projet est obligatoire.'
        if not commercial_id:
            errors['commercial'] = 'Le commercial est obligatoire.'
        if not start_date:
            errors['start_date'] = 'La date de début est obligatoire.'

        coordinateur_ids = {int(pk) for pk in request.POST.getlist('coordinateurs')}

        if errors:
            ctx = self._ctx()
            ctx.update({'errors': errors, 'post': data, 'selected_coordinateurs': [str(pk) for pk in coordinateur_ids]})
            return render(request, self.template_name, ctx)

        self.project.name = name
        self.project.client_name = client_name
        self.project.chef_de_projet_id = chef_de_projet_id
        self.project.manager_id = commercial_id
        self.project.start_date = start_date
        self.project.save(update_fields=['name', 'client_name', 'chef_de_projet', 'manager', 'start_date', 'updated_at'])
        self.project.refresh_from_db()

        meta, _ = PoseProjetMeta.objects.get_or_create(project=self.project, defaults={'region': region})
        meta.region = region
        meta.commercial_id = commercial_id
        meta.save(update_fields=['region', 'commercial', 'updated_at'])

        delegations_actives = self.project.pose_delegations.filter(actif=True)
        deja_delegues = set(delegations_actives.values_list('conducteur_delegue_id', flat=True))
        for delegation in delegations_actives.exclude(conducteur_delegue_id__in=coordinateur_ids):
            revoquer_delegation(request.user, delegation)
        nouveaux = User.objects.filter(pk__in=coordinateur_ids - deja_delegues, role='CONDUCTEUR_POSE')
        for coordinateur in nouveaux:
            try:
                deleguer_projet(request.user, self.project, coordinateur)
            except PointageWorkflowError:
                pass

        messages.success(request, f'Projet « {self.project.name} » mis à jour.')
        return redirect('projects:pose_gestion_pipeline')


class PoseGestionProjetDeleteView(LoginRequiredMixin, View):
    """Suppression d'un projet de pose — action destructive réservée DG/DC/Admin (comme la création)."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_creer_projet_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        nom = self.project.name
        self.project.delete()
        messages.success(request, f'Projet « {nom} » supprimé.')
        return redirect('projects:pose_gestion_pipeline')


def _pose_articles_kpis(articles):
    """KPIs du tableau des articles de pose — factorisé pour être recalculé identique à
    l'affichage normal (GET) et lors d'un réaffichage après erreur de validation (POST)."""
    nb_articles = len(articles)
    nb_valides = sum(1 for a in articles if a.statut_validation == Article.StatutValidation.VALIDE)
    nb_en_attente = sum(1 for a in articles if a.statut_validation == Article.StatutValidation.EN_ATTENTE)
    nb_refuses = sum(1 for a in articles if a.statut_validation == Article.StatutValidation.REFUSE)
    nb_incomplets = sum(
        1 for a in articles
        if not a.verrouille and (not a.prix_de_pose or len(a.poseurs.all()) == 0)
    )

    # Coût global de la pose = Qté × prix de la pose (idem pour le budget indicatif, qui est
    # lui aussi saisi par article, pas déjà multiplié par la quantité).
    centimes = Decimal('0.01')
    total_budget_pose = sum(
        ((a.budget_de_pose or Decimal('0')) * a.quantite for a in articles), Decimal('0'),
    ).quantize(centimes)
    montant_realise_total = sum((montant_article(a) for a in articles), Decimal('0')).quantize(centimes)
    # Dénominateur de l'avancement physique : uniquement le prix de pose (ce qui est réellement
    # posé/réglé), jamais la logistique — sinon un coût logistique élevé ferait mécaniquement
    # baisser le % d'avancement sans qu'aucun travail de pose supplémentaire n'ait été fait.
    montant_prevu_total = sum(
        ((a.prix_de_pose or Decimal('0')) * a.quantite for a in articles), Decimal('0'),
    ).quantize(centimes)
    # « Budget engagé de la pose » : prix de pose + frais de logistique pose (RG nouvelle) —
    # c'est ce total, pas montant_prevu_total, qui reflète le coût global réellement engagé.
    total_frais_logistique = sum(
        ((a.frais_logistique_pose or Decimal('0')) * a.quantite for a in articles), Decimal('0'),
    ).quantize(centimes)
    total_prix_pose = montant_prevu_total + total_frais_logistique
    depassement = (total_prix_pose - total_budget_pose) if total_prix_pose > total_budget_pose else Decimal('0.00')

    avancement_global = (
        round(float(montant_realise_total) * 100 / float(montant_prevu_total), 1)
        if montant_prevu_total else 0.0
    )

    return {
        'avancement_global': avancement_global,
        'montant_realise_total': montant_realise_total,
        'montant_prevu_total': montant_prevu_total,
        'total_budget_pose': total_budget_pose,
        'total_frais_logistique': total_frais_logistique,
        'total_prix_pose': total_prix_pose,
        'depassement': depassement,
        'nb_articles': nb_articles,
        'nb_valides': nb_valides,
        'nb_en_attente': nb_en_attente,
        'nb_refuses': nb_refuses,
        'nb_incomplets': nb_incomplets,
    }


class PoseGestionProjetTableauView(LoginRequiredMixin, View):
    """Étape 2 du workflow pose : tableau des articles (affectation poseurs + prix de pose)."""

    template_name = 'projects/pose_gestion_projet.html'

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        articles = list(
            self.project.pose_articles.all().prefetch_related('poseurs', 'decompositions__pointages')
        )
        for article in articles:
            article.verrouille = article_verrouille(article)
            article.prix_de_pose_input = None
            article.frais_logistique_input = None
            article.poseurs_ids_override = None

        return render(request, self.template_name, {
            'project': self.project,
            'meta': getattr(self.project, 'pose_meta', None),
            'articles': articles,
            'poseurs': Poseur.objects.filter(is_active=True),
            'peut_gerer_poseurs': peut_gerer_referentiel_poseurs(request.user),
            'kpis': _pose_articles_kpis(articles),
        })


class PoseGestionArticleUpdateView(LoginRequiredMixin, View):
    """Saisie du prix de pose (RG-02) + affectation des poseurs pour un article — Étape 2."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, article_pk):
        article = get_object_or_404(Article, pk=article_pk, project=self.project)

        prix = request.POST.get('prix_de_pose', '').strip()
        if prix:
            try:
                prix_decimal = Decimal(prix.replace(',', '.'))
            except InvalidOperation:
                messages.error(request, 'Prix de pose invalide.')
                return redirect('projects:pose_gestion_projet', pk=self.project.pk)

            frais_logistique = request.POST.get('frais_logistique_pose', '').strip()
            frais_logistique_decimal = None
            if frais_logistique:
                try:
                    frais_logistique_decimal = Decimal(frais_logistique.replace(',', '.'))
                except InvalidOperation:
                    messages.error(request, 'Frais de logistique pose invalides.')
                    return redirect('projects:pose_gestion_projet', pk=self.project.pk)

            definir_prix_de_pose(article, prix_decimal, frais_logistique_decimal)

        poseur_ids = request.POST.getlist('poseurs')
        Affectation.objects.filter(article=article).exclude(poseur_id__in=poseur_ids).delete()
        for pid in poseur_ids:
            Affectation.objects.get_or_create(
                article=article, poseur_id=pid, defaults={'affecte_par': request.user},
            )

        messages.success(request, f'Article {article.code} mis à jour.')
        return redirect('projects:pose_gestion_projet', pk=self.project.pk)


class PoseGestionArticlesBulkUpdateView(LoginRequiredMixin, View):
    """Saisie groupée du prix de pose (RG-02) + affectation des poseurs pour tous les
    articles d'un projet en une seule soumission, au lieu d'un enregistrement article par article."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        articles = list(
            self.project.pose_articles.all().prefetch_related('poseurs', 'decompositions__pointages')
        )
        a_appliquer = []
        erreurs = []

        for article in articles:
            article.verrouille = article_verrouille(article)
            article.prix_de_pose_input = None
            article.frais_logistique_input = None
            article.poseurs_ids_override = None

            if article.verrouille:
                continue  # article déjà payé : prix/poseurs figés, toute soumission est ignorée

            prix = request.POST.get(f'prix_de_pose_{article.pk}', '').strip()
            frais_logistique = request.POST.get(f'frais_logistique_pose_{article.pk}', '').strip()
            poseur_ids = request.POST.getlist(f'poseurs_{article.pk}')
            # on conserve exactement ce que l'utilisateur a saisi pour le réafficher tel quel
            # en cas de rejet — rien ne doit se réinitialiser (voir demande utilisateur).
            article.prix_de_pose_input = prix
            article.frais_logistique_input = frais_logistique
            article.poseurs_ids_override = poseur_ids

            if not prix:
                erreurs.append(f"{article.code} : le prix de pose est obligatoire.")
                continue
            try:
                prix_decimal = Decimal(prix.replace(',', '.'))
            except InvalidOperation:
                erreurs.append(f"{article.code} : prix de pose invalide (« {prix} »).")
                continue
            if prix_decimal <= 0:
                erreurs.append(f"{article.code} : le prix de pose doit être supérieur à zéro.")
                continue

            frais_logistique_decimal = None
            if frais_logistique:
                try:
                    frais_logistique_decimal = Decimal(frais_logistique.replace(',', '.'))
                except InvalidOperation:
                    erreurs.append(f"{article.code} : frais de logistique pose invalides (« {frais_logistique} »).")
                    continue
                if frais_logistique_decimal < 0:
                    erreurs.append(f"{article.code} : les frais de logistique pose ne peuvent pas être négatifs.")
                    continue

            if not poseur_ids:
                erreurs.append(f"{article.code} : au moins un poseur doit être affecté.")
                continue

            a_appliquer.append((article, prix_decimal, frais_logistique_decimal, poseur_ids))

        if erreurs:
            for erreur in erreurs:
                messages.error(request, erreur)
            messages.error(
                request,
                "Rien n'a été enregistré — corrigez les articles signalés ci-dessous puis réessayez. "
                "Vos saisies ont été conservées.",
            )
            return render(request, 'projects/pose_gestion_projet.html', {
                'project': self.project,
                'meta': getattr(self.project, 'pose_meta', None),
                'articles': articles,
                'poseurs': Poseur.objects.filter(is_active=True),
                'peut_gerer_poseurs': peut_gerer_referentiel_poseurs(request.user),
                'kpis': _pose_articles_kpis(articles),
            })

        for article, prix_decimal, frais_logistique_decimal, poseur_ids in a_appliquer:
            definir_prix_de_pose(article, prix_decimal, frais_logistique_decimal)
            Affectation.objects.filter(article=article).exclude(poseur_id__in=poseur_ids).delete()
            for poseur_id in poseur_ids:
                Affectation.objects.get_or_create(
                    article=article, poseur_id=poseur_id, defaults={'affecte_par': request.user},
                )

        if a_appliquer:
            messages.success(
                request,
                f"{len(a_appliquer)} article(s) enregistré(s) et soumis à validation DG.",
            )
        else:
            messages.info(request, 'Aucun article à mettre à jour (tous déjà payés/verrouillés).')
        return redirect('projects:pose_gestion_projet', pk=self.project.pk)


class PoseGestionArticlesExportXLSXView(LoginRequiredMixin, View):
    """Export .xlsx du tableau des articles d'un projet de pose — Étape 2."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        import openpyxl
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = 'Articles'

        headers = [
            'Codification', 'Désignation des articles', 'Unité', 'Qté',
            'Prix unit. HT', 'Prix total HT', 'Poseur(s) affecté(s)',
            'Budget pose (estim.)', 'Prix de pose /article', 'Frais logistique pose /article',
            'Validation',
        ]
        sheet.append(headers)
        for col in range(1, len(headers) + 1):
            sheet.cell(row=1, column=col).font = openpyxl.styles.Font(bold=True)

        articles = self.project.pose_articles.all().prefetch_related('poseurs')
        for article in articles:
            sheet.append([
                article.code, article.designation, article.unite,
                article.quantite, article.prix_unitaire_ht, article.prix_total_ht,
                ', '.join(p.nom for p in article.poseurs.all()),
                article.budget_de_pose, article.prix_de_pose, article.frais_logistique_pose,
                article.get_statut_validation_display(),
            ])

        for col in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 22

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="articles_{self.project.reference}.xlsx"'
        wb.save(response)
        return response


class PoseGestionPoseursListView(LoginRequiredMixin, View):
    """Référentiel des poseurs — vue d'ensemble et création."""

    template_name = 'projects/pose_gestion_poseurs.html'

    def dispatch(self, request, *args, **kwargs):
        if not peut_gerer_referentiel_poseurs(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        poseurs = Poseur.objects.all().order_by('nom')
        nb_actifs = sum(1 for p in poseurs if p.is_active)
        nb_sans_scan = sum(1 for p in poseurs if not p.scan_cin)
        return render(request, self.template_name, {
            'poseurs': poseurs,
            'nb_actifs': nb_actifs,
            'nb_sans_scan': nb_sans_scan,
        })


class PoseGestionPoseurCreateView(LoginRequiredMixin, View):
    """Création rapide d'un poseur — référentiel partagé du module Gestion de la pose."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_gerer_referentiel_poseurs(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request):
        nom = request.POST.get('nom', '').strip()
        cin = request.POST.get('cin', '').strip()
        scan_cin = request.FILES.get('scan_cin')
        next_url = request.POST.get('next') or reverse('projects:list')

        if not nom or not cin:
            messages.error(request, 'Nom et CIN sont obligatoires pour créer un poseur.')
        elif Poseur.objects.filter(cin=cin).exists():
            messages.error(request, f'Un poseur avec le CIN {cin} existe déjà.')
        else:
            Poseur.objects.create(nom=nom, cin=cin, scan_cin=scan_cin)
            messages.success(request, f'Poseur « {nom} » créé.')
        return redirect(next_url)


class PoseGestionPoseurScanCinUploadView(LoginRequiredMixin, View):
    """Dépôt ou remplacement du scan de la CIN pour un poseur existant."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_gerer_referentiel_poseurs(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, poseur_pk):
        poseur = get_object_or_404(Poseur, pk=poseur_pk)
        next_url = request.POST.get('next') or reverse('projects:pose_gestion_poseurs_list')
        scan_cin = request.FILES.get('scan_cin')
        if not scan_cin:
            messages.error(request, 'Merci de sélectionner un fichier à téléverser.')
        else:
            poseur.scan_cin = scan_cin
            poseur.save(update_fields=['scan_cin'])
            messages.success(request, f'Scan de la CIN mis à jour pour {poseur.nom}.')
        return redirect(next_url)


class PoseGestionValidationListView(LoginRequiredMixin, View):
    """Étape 3 du workflow pose : file de validation des prix de pose (DG), tous projets confondus."""

    template_name = 'projects/pose_gestion_validation.html'

    def dispatch(self, request, *args, **kwargs):
        if not peut_valider_prix_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        articles = Article.objects.filter(
            statut_validation=Article.StatutValidation.EN_ATTENTE,
            prix_de_pose__isnull=False,
        ).select_related('project', 'project__chef_de_projet')

        project_id = request.GET.get('project') or ''
        chef_id = request.GET.get('chef') or ''
        if project_id:
            articles = articles.filter(project_id=project_id)
        if chef_id:
            articles = articles.filter(project__chef_de_projet_id=chef_id)
        articles = articles.order_by('project__reference', 'code')

        projets = Project.objects.filter(
            pose_articles__statut_validation=Article.StatutValidation.EN_ATTENTE,
            pose_articles__prix_de_pose__isnull=False,
        ).distinct().order_by('reference')
        chefs = User.objects.filter(
            projets_diriges__pose_articles__statut_validation=Article.StatutValidation.EN_ATTENTE,
            projets_diriges__pose_articles__prix_de_pose__isnull=False,
        ).distinct().order_by('last_name')

        groupes_par_projet = {}
        ordre_projets = []
        for article in articles:
            if article.project_id not in groupes_par_projet:
                groupes_par_projet[article.project_id] = {'project': article.project, 'articles': []}
                ordre_projets.append(article.project_id)
            groupes_par_projet[article.project_id]['articles'].append(article)
        groupes = [groupes_par_projet[pid] for pid in ordre_projets]

        return render(request, self.template_name, {
            'articles': articles,
            'groupes': groupes,
            'projets': projets,
            'chefs': chefs,
            'selected_project': project_id,
            'selected_chef': chef_id,
        })


class PoseGestionValidationValiderMasseView(LoginRequiredMixin, View):
    """Validation en masse d'une sélection d'articles (ou de tous les articles d'un projet en un
    clic) — file de validation des prix de pose (DG)."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_valider_prix_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request):
        from django.utils import timezone

        article_ids = request.POST.getlist('article_ids')
        redirect_url = f"{reverse('projects:pose_gestion_validation')}?{request.POST.get('querystring', '')}"

        articles = Article.objects.filter(
            pk__in=article_ids,
            statut_validation=Article.StatutValidation.EN_ATTENTE,
            prix_de_pose__isnull=False,
        )
        nb = articles.update(
            statut_validation=Article.StatutValidation.VALIDE, motif_refus='', updated_at=timezone.now(),
        )
        if nb:
            messages.success(request, f'{nb} article(s) validé(s) — pointables.')
        else:
            messages.info(request, 'Aucun article sélectionné.')
        return redirect(redirect_url)


class PoseGestionValidationDecisionView(LoginRequiredMixin, View):
    """Décision DG sur un article : valider, réviser le prix, ou refuser (RG-03 : motif obligatoire)."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_valider_prix_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, article_pk):
        article = get_object_or_404(Article, pk=article_pk)
        action = request.POST.get('action')
        redirect_url = f"{reverse('projects:pose_gestion_validation')}?{request.POST.get('querystring', '')}"

        if action == 'valider':
            article.statut_validation = Article.StatutValidation.VALIDE
            article.motif_refus = ''
            article.save(update_fields=['statut_validation', 'motif_refus', 'updated_at'])
            messages.success(request, f'Article {article.code} validé — pointable.')
        elif action == 'reviser':
            try:
                reviser_prix_de_pose(article, request.POST.get('nouveau_prix', ''), request.user)
                messages.success(request, f'Prix de {article.code} révisé et validé — pointable.')
            except PointageError as exc:
                messages.error(request, str(exc))
        elif action == 'refuser':
            motif = request.POST.get('motif_refus', '').strip()
            if not motif:
                messages.error(request, 'Le motif de refus est obligatoire.')
                return redirect(redirect_url)
            article.statut_validation = Article.StatutValidation.REFUSE
            article.motif_refus = motif
            article.save(update_fields=['statut_validation', 'motif_refus', 'updated_at'])
            messages.success(request, f'Article {article.code} refusé.')
        else:
            messages.error(request, 'Action inconnue.')

        return redirect(redirect_url)


class PoseGestionPointageView(LoginRequiredMixin, View):
    """Étape 4 du workflow pose : pointage journalier par lots (chef de projet)."""

    template_name = 'projects/pose_gestion_pointage.html'

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        date_str = request.GET.get('date')
        try:
            date = datetime.date.fromisoformat(date_str) if date_str else datetime.date.today()
        except ValueError:
            date = datetime.date.today()

        all_articles = self.project.pose_articles.all().prefetch_related('poseurs', 'lots')

        groupes = []
        nb_non_pointables = 0
        for article in all_articles:
            if not article.poseurs.exists():
                continue
            if article.statut_validation != Article.StatutValidation.VALIDE:
                nb_non_pointables += 1
                continue

            lignes = []
            for poseur in article.poseurs.all():
                lot_ouvert = next(
                    (l for l in article.lots.all() if l.poseur_id == poseur.pk and not l.is_cloture),
                    None,
                )
                palier_min = (lot_ouvert.pourcentage + 10) if lot_ouvert else 10
                lignes.append({
                    'poseur': poseur,
                    'lot': lot_ouvert,
                    'paliers': list(range(palier_min, 101, 10)),
                })
            groupes.append({'article': article, 'lignes': lignes})

        return render(request, self.template_name, {
            'project': self.project,
            'date': date,
            'groupes': groupes,
            'nb_non_pointables': nb_non_pointables,
            'avancement_projet': pipeline_kpis(self.project)['avancement_physique'],
        })


class PoseGestionLotCreateView(LoginRequiredMixin, View):
    """Création d'un nouveau lot — Étape 4."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        date_str = request.POST.get('date') or datetime.date.today().isoformat()
        try:
            date = datetime.date.fromisoformat(date_str)
        except ValueError:
            date = datetime.date.today()

        article = get_object_or_404(Article, pk=request.POST.get('article_id'), project=self.project)
        poseur = get_object_or_404(Poseur, pk=request.POST.get('poseur_id'))

        try:
            creer_lot(
                article, poseur,
                request.POST.get('quantite', ''),
                request.POST.get('pourcentage_initial', ''),
                date, request.user,
            )
            messages.success(request, f'Lot créé pour {article.code} / {poseur.nom}.')
        except PointageError as exc:
            messages.error(request, str(exc))

        url = reverse('projects:pose_gestion_pointage', args=[self.project.pk])
        return redirect(f'{url}?date={date.isoformat()}')


class PoseGestionLotPointerView(LoginRequiredMixin, View):
    """Mise à jour (pointage) d'un lot existant — Étape 4, RG-04 à RG-07."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, lot_pk):
        lot = get_object_or_404(Lot, pk=lot_pk, article__project=self.project)
        date_str = request.POST.get('date') or datetime.date.today().isoformat()
        try:
            date = datetime.date.fromisoformat(date_str)
        except ValueError:
            date = datetime.date.today()

        try:
            enregistrer_pointage(lot, request.POST.get('pourcentage', ''), date, request.user)
            messages.success(request, f'Pointage enregistré pour {lot.article.code} / {lot.poseur.nom}.')
        except PointageError as exc:
            messages.error(request, str(exc))

        url = reverse('projects:pose_gestion_pointage', args=[self.project.pk])
        return redirect(f'{url}?date={date.isoformat()}')


def _semaine_annee_from_request(request):
    periode = request.GET.get('periode') or request.POST.get('periode')
    if periode and '-' in periode:
        try:
            s, a = periode.split('-')
            return int(s), int(a)
        except ValueError:
            pass

    today = datetime.date.today()
    iso = today.isocalendar()
    try:
        semaine = int(request.GET.get('semaine') or request.POST.get('semaine') or iso[1])
        annee = int(request.GET.get('annee') or request.POST.get('annee') or iso[0])
    except (TypeError, ValueError):
        semaine, annee = iso[1], iso[0]
    return semaine, annee


class PoseGestionReglementView(LoginRequiredMixin, View):
    """Étape 5 du workflow pose : règlement hebdomadaire par poseur (chef de projet)."""

    template_name = 'projects/pose_gestion_reglement.html'

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        semaine, annee = _semaine_annee_from_request(request)
        lignes = lignes_reglement_semaine(self.project, semaine, annee)

        semaines_disponibles = list(
            Pointage.objects.filter(project=self.project)
            .values_list('semaine', 'annee').distinct().order_by('-annee', '-semaine')
        )
        if (semaine, annee) not in semaines_disponibles:
            semaines_disponibles.insert(0, (semaine, annee))

        total_du = sum((l['du'] for l in lignes), Decimal('0'))
        total_verse = sum((l['verse'] for l in lignes), Decimal('0'))

        return render(request, self.template_name, {
            'project': self.project,
            'semaine': semaine,
            'annee': annee,
            'lignes': lignes,
            'semaines_disponibles': semaines_disponibles,
            'total_du': total_du,
            'total_verse': total_verse,
            'total_reste': total_du - total_verse,
        })


class PoseGestionPaiementCreateView(LoginRequiredMixin, View):
    """Versement partiel (ou total) à un poseur pour la semaine — Étape 5, RG-10."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        semaine, annee = _semaine_annee_from_request(request)
        poseur = get_object_or_404(Poseur, pk=request.POST.get('poseur_id'))
        date = datetime.date.today()

        try:
            verser_paiement(
                self.project, poseur, semaine, annee,
                request.POST.get('montant', ''), date, request.user,
            )
            messages.success(request, f"Versement enregistré pour {poseur.nom} — en attente d'approbation Trésorerie.")
        except PointageError as exc:
            messages.error(request, str(exc))

        url = reverse('projects:pose_gestion_reglement', args=[self.project.pk])
        return redirect(f'{url}?semaine={semaine}&annee={annee}')


class PoseGestionSolderPoseurView(LoginRequiredMixin, View):
    """Solde le reliquat exact d'un poseur pour la semaine — Étape 5."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        semaine, annee = _semaine_annee_from_request(request)
        poseur = get_object_or_404(Poseur, pk=request.POST.get('poseur_id'))
        try:
            solder_poseur_semaine(self.project, poseur, semaine, annee, datetime.date.today(), request.user)
            messages.success(
                request,
                f"Reliquat de {poseur.nom} enregistré pour la semaine {semaine}/{annee} — "
                f"en attente d'approbation Trésorerie.",
            )
        except PointageError as exc:
            messages.error(request, str(exc))

        url = reverse('projects:pose_gestion_reglement', args=[self.project.pk])
        return redirect(f'{url}?semaine={semaine}&annee={annee}')


class PoseGestionSolderProjetView(LoginRequiredMixin, View):
    """Solde tous les poseurs ayant un reste à payer, pour ce projet et cette semaine — Étape 5."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        semaine, annee = _semaine_annee_from_request(request)
        paiements = solder_projet_semaine(self.project, semaine, annee, datetime.date.today(), request.user)
        messages.success(
            request,
            f"{len(paiements)} reliquat(s) enregistré(s) pour la semaine {semaine}/{annee} — "
            f"en attente d'approbation Trésorerie.",
        )

        url = reverse('projects:pose_gestion_reglement', args=[self.project.pk])
        return redirect(f'{url}?semaine={semaine}&annee={annee}')


class PoseGestionReglementExportCSVView(LoginRequiredMixin, View):
    """Export CSV du règlement hebdomadaire — Étape 5."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        semaine, annee = _semaine_annee_from_request(request)
        lignes = lignes_reglement_semaine(self.project, semaine, annee)

        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = (
            f'attachment; filename="reglement_{self.project.reference}_S{semaine}-{annee}.csv"'
        )
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Poseur', 'CIN', 'Dû', 'Versé', 'Reste à payer', 'Statut'])
        for ligne in lignes:
            writer.writerow([
                ligne['poseur'].nom, ligne['poseur'].cin,
                ligne['du'], ligne['verse'], ligne['reste'], ligne['statut'],
            ])
        return response


class PoseGestionBonApprouverView(LoginRequiredMixin, View):
    """Approbation préalable du chef de projet sur un bon de règlement — RG-13."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        semaine, annee = _semaine_annee_from_request(request)
        poseur = get_object_or_404(Poseur, pk=request.POST.get('poseur_id'))
        try:
            approuver_bon_chef(self.project, poseur, semaine, annee, request.user)
            messages.success(request, f'Bon de règlement de {poseur.nom} approuvé — en attente de validation DG.')
        except PointageError as exc:
            messages.error(request, str(exc))

        url = reverse('projects:pose_gestion_reglement', args=[self.project.pk])
        return redirect(f'{url}?semaine={semaine}&annee={annee}')


class PoseGestionBonValiderView(LoginRequiredMixin, View):
    """Validation DG d'un bon de règlement — nécessite l'approbation préalable du chef (RG-13)."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_valider_prix_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, bon_pk):
        bon = get_object_or_404(BonReglement, pk=bon_pk)
        try:
            valider_bon_dg(bon, request.user)
            messages.success(request, f'Bon de {bon.poseur.nom} validé — éditable.')
        except PointageError as exc:
            messages.error(request, str(exc))

        return_to = request.POST.get('return_to')
        if return_to:
            return redirect(return_to)
        redirect_url = f"{reverse('projects:pose_gestion_bons_validation')}?{request.POST.get('querystring', '')}"
        return redirect(redirect_url)


class PoseGestionBonsValidationListView(LoginRequiredMixin, View):
    """File DG des bons de règlement en attente de validation (déjà approuvés par le chef), tous projets confondus."""

    template_name = 'projects/pose_gestion_bons_validation.html'

    def dispatch(self, request, *args, **kwargs):
        if not peut_valider_prix_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        bons = BonReglement.objects.filter(
            statut=BonReglement.Statut.APPROUVE_CHEF,
        ).select_related('project', 'poseur', 'approuve_par_chef').order_by('project__reference', 'poseur__nom')

        project_id = request.GET.get('project') or ''
        if project_id:
            bons = bons.filter(project_id=project_id)

        projets = Project.objects.filter(
            pose_bons_reglement__statut=BonReglement.Statut.APPROUVE_CHEF,
        ).distinct().order_by('reference')

        return render(request, self.template_name, {
            'bons': bons, 'projets': projets, 'selected_project': project_id,
        })


def _paiements_return_url(request):
    """Reconstruit l'URL de la page Paiements en conservant les filtres projet/recherche."""
    from urllib.parse import urlencode

    params = {}
    if request.POST.get('project'):
        params['project'] = request.POST['project']
    if request.POST.get('q'):
        params['q'] = request.POST['q']
    url = reverse('projects:pose_gestion_paiements')
    return f'{url}?{urlencode(params)}' if params else url


class PoseGestionPaiementsValidesListView(LoginRequiredMixin, View):
    """Vue Trésorerie : versements espèce à approuver + bons de règlement validés par la DG,
    tous projets confondus — RG-15."""

    template_name = 'projects/pose_gestion_paiements.html'

    def dispatch(self, request, *args, **kwargs):
        if not peut_editer_bon_reglement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        project_id = request.GET.get('project') or ''
        q = request.GET.get('q', '').strip()

        paiements_attente = Paiement.objects.filter(
            statut=Paiement.Statut.EN_ATTENTE,
        ).select_related('project', 'poseur', 'created_by').order_by('project__reference', '-date')

        bons = BonReglement.objects.filter(
            statut=BonReglement.Statut.VALIDE_DG,
        ).select_related('project', 'poseur', 'valide_par_dg').order_by('-annee', '-semaine', 'project__reference')

        if project_id:
            paiements_attente = paiements_attente.filter(project_id=project_id)
            bons = bons.filter(project_id=project_id)

        if q:
            recherche = (
                Q(project__name__icontains=q)
                | Q(project__reference__icontains=q)
                | Q(project__client_name__icontains=q)
            )
            paiements_attente = paiements_attente.filter(recherche)
            bons = bons.filter(recherche)

        projets = Project.objects.filter(
            Q(pose_paiements__statut=Paiement.Statut.EN_ATTENTE)
            | Q(pose_bons_reglement__statut=BonReglement.Statut.VALIDE_DG),
        ).distinct().order_by('reference')

        return render(request, self.template_name, {
            'paiements_attente': paiements_attente,
            'bons': bons, 'projets': projets, 'selected_project': project_id, 'q': q,
        })


class PoseGestionPaiementApprouverView(LoginRequiredMixin, View):
    """Approbation Trésorerie d'un versement espèce — RG-15."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_editer_bon_reglement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, paiement_pk):
        paiement = get_object_or_404(Paiement, pk=paiement_pk)
        try:
            approuver_paiement(paiement, request.user)
            messages.success(request, f'Versement de {paiement.poseur.nom} approuvé.')
        except PointageError as exc:
            messages.error(request, str(exc))

        return redirect(_paiements_return_url(request))


class PoseGestionPaiementRejeterView(LoginRequiredMixin, View):
    """Rejet Trésorerie d'un versement espèce erroné — reste tracé (jamais compté dans le versé)."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_editer_bon_reglement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, paiement_pk):
        paiement = get_object_or_404(Paiement, pk=paiement_pk)
        try:
            rejeter_paiement(paiement, request.POST.get('motif', ''), request.user)
            messages.success(request, f'Versement de {paiement.poseur.nom} rejeté.')
        except PointageError as exc:
            messages.error(request, str(exc))

        return redirect(_paiements_return_url(request))


class PoseGestionBonEditView(LoginRequiredMixin, View):
    """Visualisation du bon de règlement — aperçu dès l'approbation du chef, définitif après validation DG."""

    template_name = 'projects/pose_gestion_bon_reglement.html'

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        autorise = (
            peut_gerer_projet_pose(request.user, self.project)
            or peut_valider_prix_pose(request.user)
            or peut_editer_bon_reglement(request.user)
        )
        if not autorise:
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk, bon_pk):
        bon = get_object_or_404(BonReglement, pk=bon_pk, project=self.project)
        if bon.statut == BonReglement.Statut.EN_ATTENTE:
            messages.error(request, "Ce bon doit d'abord être approuvé par le chef de projet.")
            url = reverse('projects:pose_gestion_reglement', args=[self.project.pk])
            return redirect(f'{url}?semaine={bon.semaine}&annee={bon.annee}')

        ligne = next(
            (l for l in lignes_reglement_semaine(self.project, bon.semaine, bon.annee) if l['poseur'].pk == bon.poseur_id),
            None,
        )

        return render(request, self.template_name, {
            'project': self.project, 'bon': bon, 'ligne': ligne,
            'apercu': bon.statut != BonReglement.Statut.VALIDE_DG,
            'peut_valider': peut_valider_prix_pose(request.user),
        })


class PoseGestionPipelineView(LoginRequiredMixin, View):
    """Étape 6 du workflow pose : tableau de bord Pipeline/Pilotage — chef (ses projets) ou DG (tous)."""

    template_name = 'projects/pose_gestion_pipeline.html'

    def get(self, request):
        projets = projets_pipeline_pose(request.user).select_related('pose_meta')

        chef_id = request.GET.get('chef') or ''
        region = request.GET.get('region') or ''
        q = request.GET.get('q', '').strip()

        chefs = User.objects.filter(
            pk__in=projets.values_list('chef_de_projet_id', flat=True),
        ).order_by('last_name')
        regions = (
            PoseProjetMeta.objects.filter(project__in=projets)
            .exclude(region='').values_list('region', flat=True).distinct().order_by('region')
        )

        if chef_id:
            projets = projets.filter(chef_de_projet_id=chef_id)
        if region:
            projets = projets.filter(pose_meta__region=region)
        if q:
            projets = projets.filter(Q(name__icontains=q) | Q(reference__icontains=q) | Q(client_name__icontains=q))

        lignes = [{'project': p, 'kpis': pipeline_kpis(p)} for p in projets]

        nb_projets = len(lignes)
        avancement_moyen = (
            round(sum(l['kpis']['avancement_physique'] for l in lignes) / nb_projets, 1)
            if nb_projets else 0
        )
        budget_engage_total = sum((l['kpis']['budget_de_pose_total'] for l in lignes), Decimal('0'))
        montant_paye_total = sum((l['kpis']['montant_paye_poseurs'] for l in lignes), Decimal('0'))
        nb_depassements = sum(1 for l in lignes if l['kpis']['depassement'] > 0)

        return render(request, self.template_name, {
            'lignes': lignes,
            'chefs': chefs,
            'regions': regions,
            'selected_chef': chef_id,
            'selected_region': region,
            'q': q,
            'peut_creer_projet': peut_initier_projet_pose(request.user),
            'nb_projets': nb_projets,
            'avancement_moyen': avancement_moyen,
            'budget_engage_total': budget_engage_total,
            'montant_paye_total': montant_paye_total,
            'nb_depassements': nb_depassements,
        })


class PoseGestionCumulAttachementUpdateView(LoginRequiredMixin, View):
    """Saisie du cumul attachement hebdomadaire — Étape 6, RG-12."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        try:
            definir_cumul_attachement(self.project, request.POST.get('cumul_attachement', ''), request.user)
            messages.success(request, 'Cumul attachement mis à jour.')
        except PointageError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_pipeline')


# ---------------------------------------------------------------------------
# Décomposition d'articles + pointage par composant (workflow BROUILLON/VALIDÉ/PAYÉ)
# ---------------------------------------------------------------------------

class PoseGestionDecompositionView(LoginRequiredMixin, View):
    """Édition de la décomposition d'un article en composants pondérés — RG-01, RG-04, RG-12."""

    template_name = 'projects/pose_gestion_decomposition.html'

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        self.article = get_object_or_404(Article, pk=kwargs['article_pk'], project=self.project)
        if not peut_gerer_decomposition(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk, article_pk):
        decompositions = list(self.article.decompositions.order_by('ordre'))
        for d in decompositions:
            d.poids_pourcent = (d.poids_prix * 100).quantize(Decimal('0.01'))
        is_valid, somme, ecart = valider_somme_poids(self.article)
        disponible = max(Decimal('0'), Decimal('1') - somme)
        return render(request, self.template_name, {
            'project': self.project,
            'article': self.article,
            'decompositions': decompositions,
            'somme_poids': somme,
            'somme_valide': is_valid,
            'ecart': ecart,
            'poids_disponible_pourcent': (disponible * 100).quantize(Decimal('0.01')),
        })


class PoseGestionDecompositionCreateView(LoginRequiredMixin, View):
    """Ajout d'un composant à la décomposition d'un article."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        self.article = get_object_or_404(Article, pk=kwargs['article_pk'], project=self.project)
        if not peut_gerer_decomposition(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, article_pk):
        libelle = request.POST.get('libelle_composant', '').strip()
        poids_raw = request.POST.get('poids_prix', '').strip()
        if not libelle:
            messages.error(request, 'Le libellé du composant est obligatoire.')
            return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article_pk)
        try:
            poids = Decimal(poids_raw.replace(',', '.').replace('%', ''))
            if poids > 1:
                poids = poids / Decimal('100')
        except InvalidOperation:
            messages.error(request, 'Poids invalide.')
            return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article_pk)
        if not (Decimal('0') < poids <= Decimal('1')):
            messages.error(request, 'Le poids doit être compris entre 0 (exclu) et 100 % (inclus).')
            return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article_pk)

        somme_actuelle = self.article.decompositions.filter(actif=True).aggregate(
            s=Sum('poids_prix'))['s'] or Decimal('0')
        if somme_actuelle + poids > Decimal('1') + TOLERANCE_SOMME_POIDS:
            disponible = max(Decimal('0'), Decimal('1') - somme_actuelle)
            messages.error(
                request,
                f"Impossible d'ajouter ce composant : la décomposition est déjà à "
                f"{(somme_actuelle * 100)} % — il ne reste que {(disponible * 100)} % disponible (RG-04)."
            )
            return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article_pk)

        ordre_max = self.article.decompositions.count()
        try:
            PoseArticleDecomposition.objects.create(
                article=self.article, libelle_composant=libelle, poids_prix=poids, ordre=ordre_max + 1,
            )
            messages.success(request, f'Composant « {libelle} » ajouté.')
        except IntegrityError:
            messages.error(request, f'Un composant « {libelle} » existe déjà pour cet article.')
        return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article_pk)


class PoseGestionDecompositionUpdateView(LoginRequiredMixin, View):
    """Modification du libellé/poids/ordre/actif d'un composant — RG-04, RG-13."""

    def dispatch(self, request, *args, **kwargs):
        self.decomposition = get_object_or_404(PoseArticleDecomposition, pk=kwargs['decomposition_pk'])
        self.project = self.decomposition.article.project
        if self.project.pk != kwargs['pk']:
            raise Http404
        if not peut_gerer_decomposition(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, decomposition_pk):
        article = self.decomposition.article
        libelle = request.POST.get('libelle_composant', '').strip()
        poids_raw = request.POST.get('poids_prix', '').strip()
        ordre_raw = request.POST.get('ordre', '').strip()
        actif = request.POST.get('actif') == 'on'

        if not libelle:
            messages.error(request, 'Le libellé du composant est obligatoire.')
            return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article.pk)
        try:
            poids = Decimal(poids_raw.replace(',', '.').replace('%', ''))
            if poids > 1:
                poids = poids / Decimal('100')
        except InvalidOperation:
            messages.error(request, 'Poids invalide.')
            return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article.pk)
        if not (Decimal('0') < poids <= Decimal('1')):
            messages.error(request, 'Le poids doit être compris entre 0 (exclu) et 100 % (inclus).')
            return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article.pk)

        if actif:
            somme_autres = article.decompositions.filter(actif=True).exclude(
                pk=self.decomposition.pk).aggregate(s=Sum('poids_prix'))['s'] or Decimal('0')
            if somme_autres + poids > Decimal('1') + TOLERANCE_SOMME_POIDS:
                disponible = max(Decimal('0'), Decimal('1') - somme_autres)
                messages.error(
                    request,
                    f"Impossible d'enregistrer : les autres composants actifs totalisent déjà "
                    f"{(somme_autres * 100)} % — ce composant ne peut pas dépasser {(disponible * 100)} % (RG-04)."
                )
                return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article.pk)

        self.decomposition.libelle_composant = libelle
        self.decomposition.poids_prix = poids
        self.decomposition.ordre = int(ordre_raw) if ordre_raw.isdigit() else self.decomposition.ordre
        self.decomposition.actif = actif
        self.decomposition.save()

        is_valid, somme, ecart = valider_somme_poids(article)
        if not is_valid:
            messages.warning(request, f'Composant enregistré — attention, Σ des poids actifs = {somme * 100} % (écart {ecart * 100} %).')
        else:
            messages.success(request, 'Composant mis à jour — Σ = 100 %.')
        return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article.pk)


class PoseGestionDecompositionDeleteView(LoginRequiredMixin, View):
    """Suppression d'un composant — RG-15 : archivage (actif=False) s'il a déjà été pointé,
    suppression réelle sinon."""

    def dispatch(self, request, *args, **kwargs):
        self.decomposition = get_object_or_404(PoseArticleDecomposition, pk=kwargs['decomposition_pk'])
        self.project = self.decomposition.article.project
        if self.project.pk != kwargs['pk']:
            raise Http404
        if not peut_gerer_decomposition(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, decomposition_pk):
        article = self.decomposition.article
        libelle = self.decomposition.libelle_composant
        deja_pointe = PosePointageComposant.objects.filter(
            decomposition=self.decomposition, quantite_posee__gt=0,
        ).exists()

        if deja_pointe:
            self.decomposition.actif = False
            self.decomposition.save(update_fields=['actif', 'updated_at'])
            messages.warning(request, f'Composant « {libelle} » déjà pointé — archivé (désactivé) plutôt que supprimé (RG-15).')
        else:
            self.decomposition.delete()
            messages.success(request, f'Composant « {libelle} » supprimé.')
        return redirect('projects:pose_gestion_decomposition', pk=pk, article_pk=article.pk)


class PoseGestionPointageComposantView(LoginRequiredMixin, View):
    """Écran de pointage par composant, en quantité totale (RG-02) — équivalent RECAP DE POSE."""

    template_name = 'projects/pose_gestion_pointage_composant.html'

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not projets_visibles_pointage(request.user).filter(pk=self.project.pk).exists():
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        user = request.user
        piece = piece_en_attente(self.project)
        peut_saisir_ici = peut_saisir(user, self.project) and piece is None
        groupes = []
        montant_realise_total = Decimal('0')
        montant_prevu_total = Decimal('0')
        nb_composants_total = 0
        nb_composants_complets = 0
        poseurs_ids = set()
        for article in self.project.pose_articles.prefetch_related('decompositions__pointages', 'poseurs'):
            if not article.decompositions.filter(actif=True).exists():
                continue
            lignes = []
            for ligne in recap_pose(article):
                lignes.append(ligne)
                montant_realise_total += ligne['montant']
                nb_composants_total += 1
                if ligne['reste'] <= 0:
                    nb_composants_complets += 1
            groupes.append({'article': article, 'lignes': lignes})
            if article.prix_de_pose:
                montant_prevu_total += article.prix_de_pose * article.quantite
            poseurs_ids.update(p.pk for p in article.poseurs.all())

        montant_prevu_total = montant_prevu_total.quantize(Decimal('0.01'))
        poseurs_projet = Poseur.objects.filter(pk__in=poseurs_ids).order_by('nom')

        avancement_global = (
            round(float(montant_realise_total) * 100 / float(montant_prevu_total), 1)
            if montant_prevu_total else 0.0
        )

        historique_pieces = PosePieceComptage.objects.filter(
            project=self.project, statut=PosePieceComptage.Statut.VALIDE,
        ).select_related('saisi_par', 'valide_par').prefetch_related(
            'lignes__decomposition__article', 'lignes__decomposition__pointages',
            'lignes__decomposition__article__poseurs', 'lignes__poseur', 'lignes__medias',
        ).order_by('-date_validation')

        toutes_lignes = [
            ligne for piece_h in historique_pieces for ligne in piece_h.lignes.all()
        ]
        fiches_par_cle = {
            (f.decomposition_id, f.poseur_id): f
            for f in PosePointageComposant.objects.filter(
                decomposition_id__in={l.decomposition_id for l in toutes_lignes},
            )
        }
        for ligne in toutes_lignes:
            ligne.fiche = fiches_par_cle.get((ligne.decomposition_id, ligne.poseur_id))

        historique = []
        for piece_h in historique_pieces:
            groupes_par_article = {}
            ordre_articles = []
            for ligne in piece_h.lignes.all():
                if ligne.quantite_saisie == 0:
                    continue
                article = ligne.decomposition.article
                if article.pk not in groupes_par_article:
                    groupes_par_article[article.pk] = {'article': article, 'lignes': []}
                    ordre_articles.append(article.pk)
                groupes_par_article[article.pk]['lignes'].append(ligne)
            if not groupes_par_article:
                continue
            historique.append({
                'piece': piece_h,
                'groupes': [groupes_par_article[apk] for apk in ordre_articles],
            })

        pieces_refusees = PosePieceComptage.objects.filter(
            project=self.project, statut=PosePieceComptage.Statut.REFUSEE,
        ).select_related('saisi_par', 'refuse_par').order_by('-date_refus')

        kpis = {
            'avancement_global': avancement_global,
            'montant_realise_total': montant_realise_total,
            'montant_prevu_total': montant_prevu_total,
            'nb_composants_total': nb_composants_total,
            'nb_composants_complets': nb_composants_complets,
            'nb_pieces_validees': len(historique_pieces),
            'nb_poseurs': len(poseurs_ids),
        }

        return render(request, self.template_name, {
            'project': self.project,
            'meta': getattr(self.project, 'pose_meta', None),
            'groupes': groupes,
            'peut_saisir_ici': peut_saisir_ici,
            'piece_en_attente': piece,
            'peut_valider_piece_ici': piece is not None and peut_valider_piece(user, piece),
            'historique': historique,
            'pieces_refusees': pieces_refusees,
            'kpis': kpis,
            'poseurs_projet': poseurs_projet,
            # Coordinateur de la pose : quantités et avancement oui, coût de la pose jamais.
            'masquer_valeurs': getattr(user, 'role', None) == 'CONDUCTEUR_POSE',
        })


class PoseGestionHistoriquePointageExportView(LoginRequiredMixin, View):
    """Export .xlsx de l'historique détaillé des pièces de comptage validées d'un projet."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not projets_visibles_pointage(request.user).filter(pk=self.project.pk).exists():
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        import openpyxl
        from django.utils import timezone
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = 'Historique pointage'

        headers = [
            'Pièce', 'Article — Code', 'Article — Désignation', 'Poseur(s)', 'Composant',
            'Quantité saisie (incrément)', 'Saisi par', 'Date de création',
            'Validé par', 'Date de validation',
        ]
        sheet.append(headers)
        for col in range(1, len(headers) + 1):
            sheet.cell(row=1, column=col).font = openpyxl.styles.Font(bold=True)

        pieces = PosePieceComptage.objects.filter(
            project=self.project, statut=PosePieceComptage.Statut.VALIDE,
        ).select_related('saisi_par', 'valide_par').prefetch_related(
            'lignes__decomposition__article__poseurs', 'lignes__poseur',
        ).order_by('-date_validation')

        for piece in pieces:
            for ligne in piece.lignes.all():
                article = ligne.decomposition.article
                designation = ' '.join(article.designation.split())
                poseurs = ligne.poseur.nom if ligne.poseur else ', '.join(p.nom for p in article.poseurs.all())
                sheet.append([
                    piece.numero, article.code, designation, poseurs,
                    ligne.decomposition.libelle_composant, ligne.quantite_saisie,
                    piece.saisi_par.get_full_name() if piece.saisi_par else '',
                    timezone.localtime(piece.date_creation).strftime('%d/%m/%Y %H:%M') if piece.date_creation else '',
                    piece.valide_par.get_full_name() if piece.valide_par else '',
                    timezone.localtime(piece.date_validation).strftime('%d/%m/%Y %H:%M') if piece.date_validation else '',
                ])

        for col in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 26

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="historique_pointage_{self.project.reference}.xlsx"'
        wb.save(response)
        return response


class PoseGestionRelevePoseurView(LoginRequiredMixin, View):
    """Relevé de paiement par poseur, pour un projet donné (RG-26) : ce qui reste à payer et
    ce qui a déjà été payé, par poseur, à partir du pointage validé."""

    template_name = 'projects/pose_gestion_releve_poseur.html'

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        a_acces = peut_gerer_projet_pose(request.user, self.project) or (
            projets_visibles_pointage(request.user).filter(pk=self.project.pk).exists()
        )
        if not a_acces:
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        # Coordinateur de la pose : visibilité du relevé (qui a pointé quoi) mais jamais des
        # montants — seul le chef de projet / la supervision manipule le coût de la pose.
        masquer_valeurs = getattr(request.user, 'role', None) == 'CONDUCTEUR_POSE'
        releve = {}
        articles = self.project.pose_articles.prefetch_related(
            'poseurs', 'decompositions__pointages__poseur', 'decompositions__pointages__valide_par',
        )
        for article in articles:
            poseurs_article = list(article.poseurs.all())
            poseur_unique = poseurs_article[0] if len(poseurs_article) == 1 else None
            prix = article.prix_de_pose or Decimal('0')
            for decomposition in article.decompositions.filter(actif=True):
                for fiche in decomposition.pointages.all():
                    if fiche.quantite_posee <= 0 and fiche.quantite_payee <= 0:
                        continue
                    poseur = fiche.poseur or poseur_unique
                    if poseur is None:
                        continue

                    bucket = releve.setdefault(poseur.pk, {
                        'poseur': poseur, 'a_payer': [], 'deja_paye': [],
                        'total_a_payer': Decimal('0'), 'total_paye': Decimal('0'),
                    })

                    solde_impaye = fiche.quantite_posee - fiche.quantite_payee
                    if fiche.statut == PosePointageComposant.Statut.VALIDE and solde_impaye > 0:
                        montant_du = montant_composant(decomposition, solde_impaye, prix)
                        bucket['a_payer'].append({
                            'article': article, 'decomposition': decomposition,
                            'fiche': fiche, 'quantite': solde_impaye, 'montant': montant_du,
                        })
                        bucket['total_a_payer'] += montant_du

                    if fiche.quantite_payee > 0:
                        montant_regle = montant_composant(decomposition, fiche.quantite_payee, prix)
                        bucket['deja_paye'].append({
                            'article': article, 'decomposition': decomposition,
                            'fiche': fiche, 'quantite': fiche.quantite_payee, 'montant': montant_regle,
                        })
                        bucket['total_paye'] += montant_regle

        releves = sorted(releve.values(), key=lambda r: r['poseur'].nom)
        total_a_payer = sum((r['total_a_payer'] for r in releves), Decimal('0'))
        total_paye = sum((r['total_paye'] for r in releves), Decimal('0'))

        return render(request, self.template_name, {
            'project': self.project,
            'meta': getattr(self.project, 'pose_meta', None),
            'releves': releves,
            'total_a_payer': total_a_payer,
            'total_paye': total_paye,
            'total_general': total_a_payer + total_paye,
            'peut_marquer_paye_ici': peut_marquer_paye(request.user, self.project),
            'masquer_valeurs': masquer_valeurs,
        })


class PoseGestionRelevePoseurMarquerPayeView(LoginRequiredMixin, View):
    """Marque en lot une sélection de pointages validés comme payés (RG-26)."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_projet_pose(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        fiche_ids = request.POST.getlist('fiche_ids')
        fiches = PosePointageComposant.objects.filter(
            pk__in=fiche_ids, decomposition__article__project=self.project,
        )
        nb = 0
        erreurs = []
        for fiche in fiches:
            try:
                marquer_paye(request.user, fiche)
                nb += 1
            except PointageWorkflowError as exc:
                erreurs.append(str(exc))

        if nb:
            messages.success(request, f'{nb} pointage(s) marqué(s) comme payé(s).')
        for erreur in erreurs:
            messages.error(request, erreur)
        if not nb and not erreurs:
            messages.info(request, 'Aucun pointage sélectionné.')
        return redirect('projects:pose_gestion_releve_poseur', pk=pk)


class PoseGestionPoseurReleveGlobalView(LoginRequiredMixin, View):
    """Relevé de paiement global d'un poseur — tous ses projets confondus (RG-26), à titre
    informatif (historique, reste à payer). La création/alimentation d'une demande de paiement
    se fait depuis la file « Demandes de paiement » → Création."""

    template_name = 'projects/pose_gestion_poseur_releve.html'

    def dispatch(self, request, *args, **kwargs):
        self.poseur = get_object_or_404(Poseur, pk=kwargs['poseur_pk'])
        if not peut_gerer_referentiel_poseurs(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, poseur_pk):
        releve = releve_global_poseur(self.poseur)

        projets = {}
        for ligne in releve['a_payer'] + releve['deja_paye']:
            bucket = projets.setdefault(ligne['project'].pk, {
                'project': ligne['project'], 'a_payer': Decimal('0'), 'paye': Decimal('0'),
            })
            if ligne['fiche'].statut == PosePointageComposant.Statut.PAYE:
                bucket['paye'] += ligne['montant']
            else:
                bucket['a_payer'] += ligne['montant']
        par_projet = sorted(projets.values(), key=lambda p: p['project'].reference)

        historique_demandes = PoseDemandePaiement.objects.filter(
            lignes__poseur=self.poseur,
        ).exclude(statut__in=(
            PoseDemandePaiement.Statut.BROUILLON, PoseDemandePaiement.Statut.EN_ATTENTE,
        )).select_related('demandeur', 'valide_par').distinct().order_by('-date_creation')

        demande_en_attente = PoseDemandePaiement.objects.filter(
            lignes__poseur=self.poseur, statut=PoseDemandePaiement.Statut.EN_ATTENTE,
        ).distinct().first()

        return render(request, self.template_name, {
            'poseur': self.poseur,
            'total_a_payer': releve['total_a_payer'],
            'total_paye': releve['total_paye'],
            'par_projet': par_projet,
            'historique_demandes': historique_demandes,
            'demande_en_attente': demande_en_attente,
            'peut_creer_demande_ici': peut_creer_demande_paiement(request.user),
        })


class PoseGestionDemandesPaiementListView(LoginRequiredMixin, View):
    """File des demandes de paiement — en attente de validation par défaut, avec onglets pour
    consulter les demandes déjà validées (payées) ou retournées."""

    template_name = 'projects/pose_gestion_demandes_paiement.html'
    ONGLET_PAR_STATUT = {
        'en_attente': PoseDemandePaiement.Statut.EN_ATTENTE,
        'validees': PoseDemandePaiement.Statut.VALIDEE,
        'retournees': PoseDemandePaiement.Statut.REFUSEE,
        'toutes': None,
    }

    def dispatch(self, request, *args, **kwargs):
        if not peut_creer_demande_paiement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        onglet = request.GET.get('onglet', 'en_attente')
        if onglet not in self.ONGLET_PAR_STATUT:
            onglet = 'en_attente'
        statut = self.ONGLET_PAR_STATUT[onglet]

        numero_query = request.GET.get('numero', '').strip()
        date_debut = request.GET.get('date_debut', '').strip()
        date_fin = request.GET.get('date_fin', '').strip()

        # « Toutes » couvre les 3 statuts traitables (pas les brouillons en cours de
        # préparation par d'autres utilisateurs, hors périmètre de cette file partagée).
        demandes = PoseDemandePaiement.objects.exclude(statut=PoseDemandePaiement.Statut.BROUILLON)
        if statut is not None:
            demandes = demandes.filter(statut=statut)

        if numero_query:
            chiffres = ''.join(c for c in numero_query if c.isdigit())
            demandes = demandes.filter(pk=int(chiffres)) if chiffres else demandes.none()
        if date_debut:
            try:
                demandes = demandes.filter(date_creation__date__gte=date_debut)
            except (ValueError, ValidationError):
                pass
        if date_fin:
            try:
                demandes = demandes.filter(date_creation__date__lte=date_fin)
            except (ValueError, ValidationError):
                pass

        demandes = demandes.select_related('demandeur', 'valide_par').prefetch_related(
            'lignes__poseur', 'lignes__fiche__decomposition__article__project',
        ).order_by('-date_creation')

        for demande in demandes:
            demande.poseurs_concernes = sorted(
                {ligne.poseur for ligne in demande.lignes.all()}, key=lambda p: p.nom,
            )
            demande.projets_concernes = sorted(
                {ligne.fiche.decomposition.article.project for ligne in demande.lignes.all()},
                key=lambda p: p.reference,
            )
            if demande.statut == PoseDemandePaiement.Statut.VALIDEE:
                demande.nb_lignes_non_reglees = sum(1 for l in demande.lignes.all() if l.recu_id is None)
            if demande.statut == PoseDemandePaiement.Statut.EN_ATTENTE:
                demande.peut_valider_ici = peut_valider_demande_paiement(request.user, demande)

        compteurs = {
            cle: (
                PoseDemandePaiement.objects.exclude(statut=PoseDemandePaiement.Statut.BROUILLON).count()
                if val is None else PoseDemandePaiement.objects.filter(statut=val).count()
            )
            for cle, val in self.ONGLET_PAR_STATUT.items()
        }

        return render(request, self.template_name, {
            'demandes': demandes, 'onglet': onglet, 'compteurs': compteurs,
            'numero_query': numero_query, 'date_debut': date_debut, 'date_fin': date_fin,
        })


class PoseGestionDemandePaiementBrouillonView(LoginRequiredMixin, View):
    """Point d'entrée « Création » depuis le menu Demandes de paiement : reprend le brouillon
    en cours de l'utilisateur ou en crée un nouveau, puis va sur son formulaire de détail."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_creer_demande_paiement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        try:
            demande = obtenir_ou_creer_brouillon(request.user)
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
            return redirect('projects:pose_gestion_demandes_paiement')
        return redirect('projects:pose_gestion_demande_paiement_detail', demande_pk=demande.pk)


class PoseGestionDemandePaiementDetailView(LoginRequiredMixin, View):
    """Formulaire de la demande de paiement : en-tête (N°, demandeur, date, statut), lignes
    déjà ajoutées (retirables tant qu'« En cours »), et — tant qu'elle est modifiable —
    sélection des relevés non payés par poseur à y ajouter."""

    template_name = 'projects/pose_gestion_demande_paiement_detail.html'

    def dispatch(self, request, *args, **kwargs):
        self.demande = get_object_or_404(PoseDemandePaiement, pk=kwargs['demande_pk'])
        if not peut_creer_demande_paiement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, demande_pk):
        lignes = self.demande.lignes.select_related(
            'poseur', 'fiche__decomposition__article__project',
        ).order_by('poseur__nom')

        modifiable = peut_gerer_brouillon_demande(request.user, self.demande)

        recap_ajout = []
        if modifiable:
            deja_engagees = fiches_engagees_ids()
            for poseur in Poseur.objects.filter(is_active=True).order_by('nom'):
                releve = releve_global_poseur(poseur, exclure_fiche_ids=deja_engagees)
                if not releve['a_payer']:
                    continue
                recap_ajout.append({
                    'poseur': poseur,
                    'lignes': sorted(
                        releve['a_payer'], key=lambda l: (l['project'].reference, l['article'].code),
                    ),
                    'total': releve['total_a_payer'],
                })

        return render(request, self.template_name, {
            'demande': self.demande,
            'lignes': lignes,
            'modifiable': modifiable,
            'recap_ajout': recap_ajout,
        })


class PoseGestionDemandePaiementAjouterLignesView(LoginRequiredMixin, View):
    """Ajoute à la demande la sélection de pointages d'un poseur donné."""

    def dispatch(self, request, *args, **kwargs):
        self.demande = get_object_or_404(PoseDemandePaiement, pk=kwargs['demande_pk'])
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, demande_pk):
        poseur = get_object_or_404(Poseur, pk=request.POST.get('poseur_id'))
        fiche_ids = request.POST.getlist('fiche_ids')
        try:
            ajouter_lignes_demande(request.user, self.demande, poseur, fiche_ids)
            messages.success(request, f'{len(fiche_ids)} ligne(s) ajoutée(s) pour {poseur.nom}.')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_demande_paiement_detail', demande_pk=demande_pk)


class PoseGestionDemandePaiementLigneSupprimerView(LoginRequiredMixin, View):
    """Retire une ligne d'une demande « En cours »."""

    def dispatch(self, request, *args, **kwargs):
        self.demande = get_object_or_404(PoseDemandePaiement, pk=kwargs['demande_pk'])
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, demande_pk, ligne_pk):
        try:
            retirer_ligne_demande(request.user, self.demande, ligne_pk)
            messages.success(request, 'Ligne retirée de la demande.')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_demande_paiement_detail', demande_pk=demande_pk)


class PoseGestionDemandePaiementSoumettreView(LoginRequiredMixin, View):
    """Soumet la demande « En cours » à la DG pour validation."""

    def dispatch(self, request, *args, **kwargs):
        self.demande = get_object_or_404(PoseDemandePaiement, pk=kwargs['demande_pk'])
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, demande_pk):
        try:
            soumettre_demande_paiement(request.user, self.demande)
            messages.success(request, f'Demande {self.demande.numero} soumise à la DG pour validation.')
            return redirect('projects:pose_gestion_demandes_paiement')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
            return redirect('projects:pose_gestion_demande_paiement_detail', demande_pk=demande_pk)


class PoseGestionDemandePaiementValiderView(LoginRequiredMixin, View):
    """Validation par la DG — marque tous les pointages liés comme payés."""

    def dispatch(self, request, *args, **kwargs):
        self.demande = get_object_or_404(PoseDemandePaiement, pk=kwargs['demande_pk'])
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, demande_pk):
        try:
            valider_demande_paiement(request.user, self.demande)
            messages.success(
                request,
                f'Demande {self.demande.numero} validée — en attente de règlement par la Trésorerie.',
            )
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_demandes_paiement')


class PoseGestionDemandesPaiementValiderMasseView(LoginRequiredMixin, View):
    """Validation en masse d'une sélection de demandes de paiement en attente."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_creer_demande_paiement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request):
        demande_ids = request.POST.getlist('demande_ids')
        demandes = PoseDemandePaiement.objects.filter(
            pk__in=demande_ids, statut=PoseDemandePaiement.Statut.EN_ATTENTE,
        )
        nb, erreurs = 0, []
        for demande in demandes:
            try:
                valider_demande_paiement(request.user, demande)
                nb += 1
            except PointageWorkflowError as exc:
                erreurs.append(f'{demande.numero} : {exc}')
        if nb:
            messages.success(request, f'{nb} demande(s) de paiement validée(s).')
        for erreur in erreurs:
            messages.error(request, erreur)
        if not nb and not erreurs:
            messages.info(request, 'Aucune demande sélectionnée.')
        return redirect('projects:pose_gestion_demandes_paiement')


class PoseGestionDemandePaiementRefuserView(LoginRequiredMixin, View):
    """Refus d'une demande de paiement par la DG."""

    def dispatch(self, request, *args, **kwargs):
        self.demande = get_object_or_404(PoseDemandePaiement, pk=kwargs['demande_pk'])
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, demande_pk):
        try:
            refuser_demande_paiement(request.user, self.demande, request.POST.get('motif', ''))
            messages.success(request, f'Demande {self.demande.numero} refusée.')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_demandes_paiement')


class PoseGestionDemandePaiementEditionView(LoginRequiredMixin, View):
    """Édition papier d'une demande de paiement validée — à faire signer par le(s) chef(s) de
    projet concerné(s) puis déposer auprès de la trésorerie."""

    template_name = 'projects/pose_gestion_demande_paiement_edition.html'

    def dispatch(self, request, *args, **kwargs):
        self.demande = get_object_or_404(PoseDemandePaiement, pk=kwargs['demande_pk'])
        if not peut_creer_demande_paiement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, demande_pk):
        if self.demande.statut != PoseDemandePaiement.Statut.VALIDEE:
            messages.error(request, "L'édition papier n'est disponible qu'une fois la demande validée.")
            return redirect('projects:pose_gestion_demande_paiement_detail', demande_pk=demande_pk)

        lignes = self.demande.lignes.select_related(
            'poseur', 'fiche__decomposition__article__project__chef_de_projet',
        ).order_by('poseur__nom')

        groupes_par_poseur = {}
        ordre_poseurs = []
        for ligne in lignes:
            if ligne.poseur_id not in groupes_par_poseur:
                groupes_par_poseur[ligne.poseur_id] = {
                    'poseur': ligne.poseur, 'lignes': [], 'sous_total': Decimal('0'),
                }
                ordre_poseurs.append(ligne.poseur_id)
            groupes_par_poseur[ligne.poseur_id]['lignes'].append(ligne)
            groupes_par_poseur[ligne.poseur_id]['sous_total'] += ligne.montant
        groupes = [groupes_par_poseur[pid] for pid in ordre_poseurs]

        chefs_concernes = sorted(
            {
                ligne.fiche.decomposition.article.project.chef_de_projet
                for ligne in lignes if ligne.fiche.decomposition.article.project.chef_de_projet_id
            },
            key=lambda u: u.get_full_name() or u.username,
        )

        qr_code_data_uri = self._generer_qr_code(request)

        return render(request, self.template_name, {
            'demande': self.demande,
            'groupes': groupes,
            'chefs_concernes': chefs_concernes,
            'qr_code_data_uri': qr_code_data_uri,
        })

    def _generer_qr_code(self, request):
        """QR code identifiant le document — encode l'URL de vérification (détail de la
        demande) pour un contrôle rapide côté trésorerie."""
        import base64
        import io

        import qrcode

        url = request.build_absolute_uri(
            reverse('projects:pose_gestion_demande_paiement_detail', args=[self.demande.pk])
        )
        img = qrcode.make(url, border=2)
        buffer = io.BytesIO()
        img.save(buffer, format='PNG')
        encoded = base64.b64encode(buffer.getvalue()).decode('ascii')
        return f'data:image/png;base64,{encoded}'


class PoseGestionDemandePaiementEnregistrerImpressionView(LoginRequiredMixin, View):
    """Incrémente le compteur d'impressions d'une demande de paiement — appelé côté client au
    déclenchement réel de l'impression (événement navigateur `beforeprint`), pas à la simple
    ouverture de la page."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_creer_demande_paiement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, demande_pk):
        PoseDemandePaiement.objects.filter(pk=demande_pk).update(nb_impressions=F('nb_impressions') + 1)
        return HttpResponse(status=204)


class PoseGestionRecusPaiementListView(LoginRequiredMixin, View):
    """File Trésorerie : lignes de demandes Validées pas encore réglées (regroupées par
    poseur, sélectionnables pour émettre un reçu), et historique des reçus déjà émis."""

    template_name = 'projects/pose_gestion_recus_paiement.html'

    def dispatch(self, request, *args, **kwargs):
        if not peut_gerer_tresorerie_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        onglet = request.GET.get('onglet', 'a_regler')
        if onglet not in ('a_regler', 'historique'):
            onglet = 'a_regler'

        groupes = []
        if onglet == 'a_regler':
            groupes_par_poseur = {}
            ordre = []
            for ligne in lignes_validees_non_payees():
                if ligne.poseur_id not in groupes_par_poseur:
                    groupes_par_poseur[ligne.poseur_id] = {'poseur': ligne.poseur, 'lignes': [], 'total': Decimal('0')}
                    ordre.append(ligne.poseur_id)
                groupes_par_poseur[ligne.poseur_id]['lignes'].append(ligne)
                groupes_par_poseur[ligne.poseur_id]['total'] += ligne.montant
            groupes = [groupes_par_poseur[pid] for pid in ordre]

        numero_query = request.GET.get('numero', '').strip()
        date_debut = request.GET.get('date_debut', '').strip()
        date_fin = request.GET.get('date_fin', '').strip()
        recus = []
        if onglet == 'historique':
            recus = PoseRecuPaiement.objects.select_related('cree_par', 'remis_par').prefetch_related(
                'lignes_payees__poseur',
            ).order_by('-created_at')
            if numero_query:
                chiffres = ''.join(c for c in numero_query if c.isdigit())
                recus = recus.filter(pk=int(chiffres)) if chiffres else recus.none()
            if date_debut:
                try:
                    recus = recus.filter(date_recu__gte=date_debut)
                except (ValueError, ValidationError):
                    pass
            if date_fin:
                try:
                    recus = recus.filter(date_recu__lte=date_fin)
                except (ValueError, ValidationError):
                    pass
            for recu in recus:
                recu.poseurs_concernes = sorted(
                    {ligne.poseur for ligne in recu.lignes_payees.all()}, key=lambda p: p.nom,
                )

        return render(request, self.template_name, {
            'onglet': onglet,
            'groupes': groupes,
            'recus': recus,
            'nb_a_regler': lignes_validees_non_payees().count(),
            'nb_historique': PoseRecuPaiement.objects.count(),
            'numero_query': numero_query, 'date_debut': date_debut, 'date_fin': date_fin,
        })


class PoseGestionRecuPaiementCreerView(LoginRequiredMixin, View):
    """Émet un reçu de paiement Trésorerie pour la sélection de lignes soumise (un ou
    plusieurs poseurs à la fois)."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_gerer_tresorerie_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request):
        ligne_ids = request.POST.getlist('ligne_ids')
        try:
            recu = creer_recu_paiement(request.user, ligne_ids)
            messages.success(request, f'Reçu {recu.numero} émis — {recu.montant_total} MAD réglés.')
            return redirect('projects:pose_gestion_recu_paiement_detail', recu_pk=recu.pk)
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
            return redirect('projects:pose_gestion_recus_paiement')


def _grouper_lignes_recu_par_poseur(recu):
    lignes = recu.lignes_payees.select_related(
        'poseur', 'fiche__decomposition__article__project__chef_de_projet', 'demande',
    ).order_by('poseur__nom')
    groupes_par_poseur = {}
    ordre = []
    for ligne in lignes:
        if ligne.poseur_id not in groupes_par_poseur:
            groupes_par_poseur[ligne.poseur_id] = {'poseur': ligne.poseur, 'lignes': [], 'sous_total': Decimal('0')}
            ordre.append(ligne.poseur_id)
        groupes_par_poseur[ligne.poseur_id]['lignes'].append(ligne)
        groupes_par_poseur[ligne.poseur_id]['sous_total'] += ligne.montant
    return [groupes_par_poseur[pid] for pid in ordre]


class PoseGestionRecuPaiementDetailView(LoginRequiredMixin, View):
    """Écran de détail d'un reçu de paiement déjà émis."""

    template_name = 'projects/pose_gestion_recu_paiement_detail.html'

    def dispatch(self, request, *args, **kwargs):
        self.recu = get_object_or_404(PoseRecuPaiement, pk=kwargs['recu_pk'])
        if not peut_gerer_tresorerie_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, recu_pk):
        return render(request, self.template_name, {
            'recu': self.recu,
            'groupes': _grouper_lignes_recu_par_poseur(self.recu),
        })


class PoseGestionRecuPaiementEditionView(LoginRequiredMixin, View):
    """Édition papier du reçu de paiement — signé par la Trésorerie (remis par) et par le(s)
    chef(s) de projet concerné(s) ou, à défaut, le Responsable de la pose (DG Pose) — jamais
    par les poseurs eux-mêmes, qui n'ont pas de compte utilisateur signataire."""

    template_name = 'projects/pose_gestion_recu_paiement_edition.html'

    def dispatch(self, request, *args, **kwargs):
        self.recu = get_object_or_404(PoseRecuPaiement, pk=kwargs['recu_pk'])
        if not peut_gerer_tresorerie_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, recu_pk):
        lignes = self.recu.lignes_payees.select_related(
            'fiche__decomposition__article__project__chef_de_projet',
        )
        chefs_concernes = sorted(
            {
                ligne.fiche.decomposition.article.project.chef_de_projet
                for ligne in lignes if ligne.fiche.decomposition.article.project.chef_de_projet_id
            },
            key=lambda u: u.get_full_name() or u.username,
        )
        if not chefs_concernes:
            dg_pose = User.objects.filter(role='DG_POSE', is_active=True).order_by('first_name').first()
            if dg_pose:
                chefs_concernes = [dg_pose]

        return render(request, self.template_name, {
            'recu': self.recu,
            'groupes': _grouper_lignes_recu_par_poseur(self.recu),
            'chefs_concernes': chefs_concernes,
            'qr_code_data_uri': self._generer_qr_code(request),
        })

    def _generer_qr_code(self, request):
        import base64
        import io

        import qrcode

        url = request.build_absolute_uri(
            reverse('projects:pose_gestion_recu_paiement_detail', args=[self.recu.pk])
        )
        img = qrcode.make(url, border=2)
        buffer = io.BytesIO()
        img.save(buffer, format='PNG')
        encoded = base64.b64encode(buffer.getvalue()).decode('ascii')
        return f'data:image/png;base64,{encoded}'


class PoseGestionRecuPaiementEnregistrerImpressionView(LoginRequiredMixin, View):
    """Incrémente le compteur d'impressions d'un reçu de paiement (événement `beforeprint`)."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_gerer_tresorerie_pose(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, recu_pk):
        PoseRecuPaiement.objects.filter(pk=recu_pk).update(nb_impressions=F('nb_impressions') + 1)
        return HttpResponse(status=204)


class PoseGestionPointageComposantSaisirView(LoginRequiredMixin, View):
    """Saisie/mise à jour d'une fiche en BROUILLON — RG-02, RG-06."""

    def dispatch(self, request, *args, **kwargs):
        self.decomposition = get_object_or_404(PoseArticleDecomposition, pk=kwargs['decomposition_pk'])
        self.project = self.decomposition.article.project
        if self.project.pk != kwargs['pk']:
            raise Http404
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, decomposition_pk):
        try:
            enregistrer_quantite(request.user, self.decomposition, request.POST.get('quantite_posee', ''))
            messages.success(request, f'Pointage enregistré pour « {self.decomposition.libelle_composant} ».')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_pointage_composant', pk=pk)


class PoseGestionPointageComposantBulkSaisirView(LoginRequiredMixin, View):
    """Saisie groupée des quantités posées de tous les composants d'un projet, en une
    seule soumission, au lieu d'un enregistrement composant par composant."""

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not projets_visibles_pointage(request.user).filter(pk=self.project.pk).exists():
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk):
        decompositions = PoseArticleDecomposition.objects.filter(
            article__project=self.project, actif=True,
        ).select_related('article').prefetch_related('article__poseurs')

        quantites = {}
        medias = {}
        for decomposition in decompositions:
            poseurs = list(decomposition.article.poseurs.all())
            if len(poseurs) >= 2:
                for poseur in poseurs:
                    cle = f'{decomposition.pk}:{poseur.pk}'
                    valeur = request.POST.get(f'quantite_{cle}', '').strip()
                    if valeur:
                        quantites[cle] = valeur
                    fichiers = request.FILES.getlist(f'media_{cle}')
                    if fichiers:
                        medias[cle] = fichiers
            else:
                cle = str(decomposition.pk)
                valeur = request.POST.get(f'quantite_{cle}', '').strip()
                if valeur:
                    quantites[cle] = valeur
                fichiers = request.FILES.getlist(f'media_{cle}')
                if fichiers:
                    medias[cle] = fichiers

        if not quantites:
            messages.info(request, 'Aucune quantité saisie.')
            return redirect('projects:pose_gestion_pointage_composant', pk=pk)

        try:
            piece = enregistrer_pointages_groupe(
                request.user, self.project, quantites, request=request, medias=medias,
            )
            messages.success(
                request,
                f'Pièce de comptage {piece.numero} créée avec {len(quantites)} pointage(s) '
                f'— en attente de validation.',
            )
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))

        return redirect('projects:pose_gestion_pointage_composant', pk=pk)


class PoseGestionPieceComptageValiderView(LoginRequiredMixin, View):
    """Validation d'une pièce de comptage (tout le lot saisi en une fois) — RG-18, RG-25.
    Débloque le projet pour une nouvelle saisie une fois validée."""

    def dispatch(self, request, *args, **kwargs):
        self.piece = get_object_or_404(PosePieceComptage, pk=kwargs['piece_pk'])
        self.project = self.piece.project
        if self.project.pk != kwargs['pk']:
            raise Http404
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, piece_pk):
        try:
            valider_piece(request.user, self.piece)
            messages.success(request, f'Pièce de comptage {self.piece.numero} validée.')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_pointage_composant', pk=pk)


class PoseGestionPieceComptageRefuserView(LoginRequiredMixin, View):
    """Rejet définitif ou demande de révision d'une pièce de comptage — RG-25. Annule les
    quantités qu'elle portait et débloque immédiatement le projet."""

    def dispatch(self, request, *args, **kwargs):
        self.piece = get_object_or_404(PosePieceComptage, pk=kwargs['piece_pk'])
        self.project = self.piece.project
        if self.project.pk != kwargs['pk']:
            raise Http404
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, piece_pk):
        type_refus = request.POST.get('type_refus', '').strip()
        motif_categorie = request.POST.get('motif_categorie', '').strip()
        motif = request.POST.get('motif', '').strip()
        try:
            refuser_piece(request.user, self.piece, type_refus, motif_categorie, motif)
            if type_refus == PosePieceComptage.TypeRefus.REJET:
                messages.success(request, f'Pièce de comptage {self.piece.numero} rejetée — quantités annulées, projet débloqué.')
            else:
                messages.success(request, f'Révision demandée sur la pièce {self.piece.numero} — quantités annulées, le chef de projet peut resoumettre.')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_pointage_composant', pk=pk)


class PoseGestionPiecesComptageEnAttenteListView(LoginRequiredMixin, View):
    """File centralisée (Responsable de la pose) de toutes les pièces de comptage en attente
    de validation, tous projets confondus — avec validation en masse."""

    template_name = 'projects/pose_gestion_pieces_comptage_en_attente.html'

    def dispatch(self, request, *args, **kwargs):
        if not peut_creer_demande_paiement(request.user):  # même périmètre : supervision pose
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        pieces = PosePieceComptage.objects.filter(
            statut=PosePieceComptage.Statut.EN_ATTENTE,
        ).select_related('project', 'saisi_par').prefetch_related(
            'lignes__decomposition__article', 'lignes__poseur',
        ).order_by('date_creation')

        for piece in pieces:
            piece.peut_valider_ici = peut_valider_piece(request.user, piece)
            piece.nb_lignes = piece.lignes.count()

        return render(request, self.template_name, {'pieces': pieces})


class PoseGestionPiecesComptageValiderMasseView(LoginRequiredMixin, View):
    """Validation en masse d'une sélection de pièces de comptage."""

    def dispatch(self, request, *args, **kwargs):
        if not peut_creer_demande_paiement(request.user):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request):
        piece_ids = request.POST.getlist('piece_ids')
        pieces = PosePieceComptage.objects.filter(
            pk__in=piece_ids, statut=PosePieceComptage.Statut.EN_ATTENTE,
        )
        nb, erreurs = 0, []
        for piece in pieces:
            try:
                valider_piece(request.user, piece)
                nb += 1
            except PointageWorkflowError as exc:
                erreurs.append(f'{piece.numero} : {exc}')
        if nb:
            messages.success(request, f'{nb} pièce(s) de comptage validée(s).')
        for erreur in erreurs:
            messages.error(request, erreur)
        if not nb and not erreurs:
            messages.info(request, 'Aucune pièce sélectionnée.')
        return redirect('projects:pose_gestion_pieces_comptage_en_attente')


class PoseGestionDemandeModificationCreateView(LoginRequiredMixin, View):
    """Demande de modification (fiche VALIDÉ → Directeur) ou de dérogation (fiche PAYÉ → DG/Admin) — RG-20, RG-21."""

    def dispatch(self, request, *args, **kwargs):
        self.fiche = get_object_or_404(PosePointageComposant, pk=kwargs['fiche_pk'])
        self.project = self.fiche.decomposition.article.project
        if self.project.pk != kwargs['pk']:
            raise Http404
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, fiche_pk):
        nouvelle_quantite = request.POST.get('nouvelle_quantite', '').strip() or None
        try:
            demander_modification(
                request.user, self.fiche, request.POST.get('motif', ''), nouvelle_quantite=nouvelle_quantite,
            )
            messages.success(request, 'Demande envoyée — en attente de traitement.')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_pointage_composant', pk=pk)


class PoseGestionDemandesModificationListView(LoginRequiredMixin, View):
    """File d'attente unifiée : modifications post-validation (Administrateur) + dérogations
    paiement (Directeur/DG_POSE/Admin)."""

    template_name = 'projects/pose_gestion_demandes_modification.html'

    def dispatch(self, request, *args, **kwargs):
        if not (peut_approuver_modification(request.user) or peut_accorder_derogation(request.user)):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        demandes = PoseDemandeModification.objects.filter(
            statut=PoseDemandeModification.Statut.EN_ATTENTE,
        ).select_related(
            'fiche__decomposition__article__project', 'demandeur',
        ).order_by('-created_at')
        for demande in demandes:
            if demande.type_demande == PoseDemandeModification.TypeDemande.VALIDATION_DIRECTEUR:
                demande.peut_traiter_ici = peut_approuver_modification(request.user)
            else:
                demande.peut_traiter_ici = peut_accorder_derogation(request.user)
        return render(request, self.template_name, {'demandes': demandes})


class PoseGestionDemandeApprouverView(LoginRequiredMixin, View):
    """Approuve une demande — dispatché sur approuver_modification ou accorder_derogation selon le type."""

    def dispatch(self, request, *args, **kwargs):
        self.demande = get_object_or_404(PoseDemandeModification, pk=kwargs['demande_pk'])
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, demande_pk):
        try:
            if self.demande.type_demande == PoseDemandeModification.TypeDemande.VALIDATION_DIRECTEUR:
                approuver_modification(request.user, self.demande)
            else:
                accorder_derogation(request.user, self.demande)
            messages.success(request, 'Demande approuvée.')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_demandes_modification')


class PoseGestionDemandeRefuserView(LoginRequiredMixin, View):
    """Refuse une demande de modification ou de dérogation."""

    def dispatch(self, request, *args, **kwargs):
        self.demande = get_object_or_404(PoseDemandeModification, pk=kwargs['demande_pk'])
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, demande_pk):
        try:
            refuser_demande(request.user, self.demande, request.POST.get('motif_refus', ''))
            messages.success(request, 'Demande refusée.')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_demandes_modification')


class PoseGestionDemandesModificationValiderMasseView(LoginRequiredMixin, View):
    """Validation en masse (« Responsable de la pose ») de plusieurs demandes de modification/
    dérogation sélectionnées à la fois."""

    def dispatch(self, request, *args, **kwargs):
        if not (peut_approuver_modification(request.user) or peut_accorder_derogation(request.user)):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request):
        demande_ids = request.POST.getlist('demande_ids')
        demandes = PoseDemandeModification.objects.filter(
            pk__in=demande_ids, statut=PoseDemandeModification.Statut.EN_ATTENTE,
        )
        nb, erreurs = 0, []
        for demande in demandes:
            try:
                if demande.type_demande == PoseDemandeModification.TypeDemande.VALIDATION_DIRECTEUR:
                    approuver_modification(request.user, demande)
                else:
                    accorder_derogation(request.user, demande)
                nb += 1
            except PointageWorkflowError as exc:
                erreurs.append(str(exc))
        if nb:
            messages.success(request, f'{nb} demande(s) approuvée(s).')
        for erreur in erreurs:
            messages.error(request, erreur)
        if not nb and not erreurs:
            messages.info(request, 'Aucune demande sélectionnée.')
        return redirect('projects:pose_gestion_demandes_modification')


class PoseGestionDelegationView(LoginRequiredMixin, View):
    """Délégation d'un projet à un conducteur de travaux — RG-16, RG-17."""

    template_name = 'projects/pose_gestion_delegations.html'

    def dispatch(self, request, *args, **kwargs):
        self.project = get_object_or_404(Project, pk=kwargs['pk'])
        if not peut_gerer_delegation(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        conducteurs = User.objects.filter(role='CONDUCTEUR_POSE', is_active_employee=True)
        return render(request, self.template_name, {
            'project': self.project,
            'delegations': self.project.pose_delegations.select_related('conducteur_delegue', 'chef_delegant'),
            'conducteurs': conducteurs,
        })

    def post(self, request, pk):
        conducteur = get_object_or_404(User, pk=request.POST.get('conducteur_id'), role='CONDUCTEUR_POSE')
        try:
            deleguer_projet(
                request.user, self.project, conducteur,
                date_debut=request.POST.get('date_debut') or None,
                date_fin=request.POST.get('date_fin') or None,
            )
            messages.success(request, f'Projet délégué à {conducteur.get_full_name() or conducteur.username}.')
        except PointageWorkflowError as exc:
            messages.error(request, str(exc))
        return redirect('projects:pose_gestion_delegations', pk=pk)


class PoseGestionDelegationRevoquerView(LoginRequiredMixin, View):
    """Révocation d'une délégation."""

    def dispatch(self, request, *args, **kwargs):
        self.delegation = get_object_or_404(PoseDelegation, pk=kwargs['delegation_pk'])
        self.project = self.delegation.project
        if self.project.pk != kwargs['pk']:
            raise Http404
        if not peut_gerer_delegation(request.user, self.project):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, pk, delegation_pk):
        revoquer_delegation(request.user, self.delegation)
        messages.success(request, 'Délégation révoquée.')
        return redirect('projects:pose_gestion_delegations', pk=pk)


class PoseGestionMesProjetsPointageView(LoginRequiredMixin, View):
    """Point d'entrée pour le pointage par composant — RG-17 : liste des projets où
    l'utilisateur est chef, délégué, ou en supervision (Directeur/DG/Admin)."""

    template_name = 'projects/pose_gestion_mes_projets_pointage.html'

    def get(self, request):
        projets = projets_visibles_pointage(request.user).select_related('pose_meta').order_by('-created_at')
        return render(request, self.template_name, {'projets': projets})


# ---------------------------------------------------------------------------
# Planning / Gantt (séquençage standard Bois&Co)
# ---------------------------------------------------------------------------

class ProjectGanttView(LoginRequiredMixin, View):
    template_name = 'projects/project_gantt.html'

    def get(self, request, pk):
        project = get_object_or_404(Project, pk=pk)
        phases = list(project.phases.order_by('order'))

        # Regroupement par partie pour l'affichage
        parties = {
            'A': {'label': PARTIE_LABELS['A'], 'color': 'foret',  'phases': []},
            'B': {'label': PARTIE_LABELS['B'], 'color': 'blue',   'phases': []},
            'C': {'label': PARTIE_LABELS['C'], 'color': 'amber',  'phases': []},
        }
        for ph in phases:
            key = ph.partie if ph.partie in parties else 'A'
            parties[key]['phases'].append(ph)

        return render(request, self.template_name, {
            'project':    project,
            'phases':     phases,
            'parties':    parties,
            'step_range': range(1, 19),
        })


class GanttPhaseUpdateView(LoginRequiredMixin, View):
    """Met à jour les dates planifiées et le statut d'une phase du Gantt.
    Réservé aux rôles responsables de la phase (+ DIRECTEUR/MANAGER/admin).
    """

    def post(self, request, pk, phase_pk):
        project = get_object_or_404(Project, pk=pk)
        phase   = get_object_or_404(Phase, pk=phase_pk, project=project)
        if not _can_comment_phase(request.user, phase):
            from django.http import HttpResponseForbidden
            return HttpResponseForbidden('Votre rôle ne vous permet pas de modifier cette étape.')

        phase.planned_start  = request.POST.get('planned_start')  or None
        phase.planned_end    = request.POST.get('planned_end')    or None
        phase.is_completed   = request.POST.get('is_completed') == '1'
        phase.is_active      = request.POST.get('is_active')    == '1'
        try:
            phase.progress = max(0, min(100, int(request.POST.get('progress', phase.progress))))
        except (ValueError, TypeError):
            pass
        # Forcer 100% si marquée terminée
        if phase.is_completed:
            phase.progress = 100
        phase.save(update_fields=['planned_start', 'planned_end', 'is_completed', 'is_active', 'progress'])

        messages.success(request, f'Étape "{phase.name}" mise à jour.')
        # Rediriger vers le projet si appelé depuis le détail, sinon vers le gantt
        next_url = request.POST.get('next') or request.META.get('HTTP_REFERER', '')
        if 'planning' not in next_url:
            return redirect('projects:detail', pk=project.pk)
        return redirect('projects:gantt', pk=project.pk)


# ---------------------------------------------------------------------------
# Vues Django — Chef de projet (formulaires HTML)
# ---------------------------------------------------------------------------

class ChefDesignFormView(LoginRequiredMixin, View):
    """POST depuis le formulaire de la fiche projet — désigne un chef de projet."""

    def post(self, request, pk):
        if not (request.user.is_superuser or request.user.role in ('DIRECTEUR', 'MANAGER')):
            from django.http import HttpResponseForbidden
            return HttpResponseForbidden()
        project = get_object_or_404(Project, pk=pk)
        chef_id = request.POST.get('chef_id')
        if chef_id:
            try:
                chef = User.objects.get(pk=chef_id)
                desig_chef_de_projet(project, chef, request.user)
                messages.success(request, f'{chef.get_full_name()} désigné chef de projet.')
            except User.DoesNotExist:
                messages.error(request, 'Utilisateur introuvable.')
        return redirect('projects:detail', pk=pk)


class ChefValidateFormView(LoginRequiredMixin, View):
    """POST depuis la fiche projet — valide la désignation du chef de projet."""

    def post(self, request, pk):
        if not (request.user.is_superuser or request.user.role in ('DIRECTEUR', 'MANAGER')):
            from django.http import HttpResponseForbidden
            return HttpResponseForbidden()
        project = get_object_or_404(Project, pk=pk)
        try:
            validate_chef(project, request.user)
            messages.success(request, 'Désignation validée.')
        except ValueError as e:
            messages.error(request, str(e))
        return redirect('projects:detail', pk=pk)


# ---------------------------------------------------------------------------
# Vues Django — Commentaires ProjectStep (formulaires HTML)
# ---------------------------------------------------------------------------

class ProjectStepCommentFormView(LoginRequiredMixin, View):
    """POST depuis la fiche projet — ajoute un commentaire sur une étape projet."""

    def post(self, request, pk):
        from .models import StepComment
        step = get_object_or_404(ProjectStep, pk=pk)
        can = (
            request.user.is_superuser
            or request.user.role in ('DIRECTEUR', 'MANAGER')
            or request.user.role in (step.responsables_roles or [])
        )
        if not can:
            from django.http import HttpResponseForbidden
            return HttpResponseForbidden('Vous ne pouvez pas commenter cette étape.')
        text = request.POST.get('text', '').strip()
        attachment = request.FILES.get('attachment') or None
        if text or attachment:
            StepComment.objects.create(step=step, author=request.user, text=text, attachment=attachment)
        return redirect('projects:detail', pk=step.project_id)


class ProjectStepCompleteFormView(LoginRequiredMixin, View):
    """POST depuis la fiche projet — marque une étape projet comme terminée."""

    def post(self, request, pk):
        step = get_object_or_404(ProjectStep, pk=pk)
        can = (
            request.user.is_superuser
            or request.user.role in ('DIRECTEUR', 'MANAGER')
            or step.project.chef_de_projet_id == request.user.pk
        )
        if not can:
            from django.http import HttpResponseForbidden
            return HttpResponseForbidden()
        complete_step(step, request.user)
        messages.success(request, f'Étape « {step.name} » marquée terminée.')
        return redirect('projects:detail', pk=step.project_id)


# ---------------------------------------------------------------------------
# API DRF — Chef de projet
# ---------------------------------------------------------------------------

class DesignChefView(APIView):
    """POST /api/projets/<pk>/chef/ — désigne un chef de projet."""
    permission_classes = [IsAuthenticated, IsManagerOrAbove]

    def post(self, request, pk):
        project = get_object_or_404(Project, pk=pk)
        chef_id = request.data.get('chef_id')
        if not chef_id:
            return Response({'detail': 'chef_id requis.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            chef = User.objects.get(pk=chef_id)
        except User.DoesNotExist:
            return Response({'detail': 'Utilisateur introuvable.'}, status=status.HTTP_404_NOT_FOUND)
        desig_chef_de_projet(project, chef, request.user)
        return Response({
            'chef_id': chef.pk,
            'chef_name': chef.get_full_name() or chef.username,
            'designe_le': project.chef_designe_le,
            'valide': project.chef_valide_par_admin,
        })


class ValidateChefView(APIView):
    """POST /api/projets/<pk>/chef/valider/ — valide la désignation (DIRECTEUR/MANAGER)."""
    permission_classes = [IsAuthenticated, IsManagerOrAbove]

    def post(self, request, pk):
        project = get_object_or_404(Project, pk=pk)
        try:
            validate_chef(project, request.user)
        except ValueError as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        return Response({'valide': True, 'valide_le': project.chef_valide_le})


# ---------------------------------------------------------------------------
# API DRF — Étapes projet
# ---------------------------------------------------------------------------

class ProjectStepListView(generics.ListCreateAPIView):
    """GET/POST /api/projets/<pk>/etapes/"""
    serializer_class = ProjectStepSerializer
    permission_classes = [IsAuthenticated, IsProjectMember]

    def get_queryset(self):
        return ProjectStep.objects.filter(project_id=self.kwargs['pk']).order_by('order')

    def perform_create(self, serializer):
        project = get_object_or_404(Project, pk=self.kwargs['pk'])
        serializer.save(project=project)


class ProjectStepCompleteView(APIView):
    """POST /api/etapes/<pk>/terminer/ — marque une étape comme terminée."""
    permission_classes = [IsAuthenticated, IsChefDeProjet]

    def get_project_pk(self):
        step = get_object_or_404(ProjectStep, pk=self.kwargs['pk'])
        return step.project_id

    def post(self, request, pk):
        step = get_object_or_404(ProjectStep, pk=pk)
        # Check permission via project
        if not (
            request.user.is_superuser
            or request.user.role in ('DIRECTEUR', 'MANAGER')
            or step.project.chef_de_projet_id == request.user.pk
        ):
            return Response({'detail': 'Permission refusée.'}, status=status.HTTP_403_FORBIDDEN)
        complete_step(step, request.user)
        return Response({'is_completed': True, 'completed_at': step.completed_at})


# ---------------------------------------------------------------------------
# API DRF — Commentaires étapes
# ---------------------------------------------------------------------------

class StepCommentListCreateView(generics.ListCreateAPIView):
    """GET/POST /api/etapes/<pk>/commentaires/
    GET: accessible à tous les membres du projet.
    POST: réservé aux rôles dans step.responsables_roles (+ DIRECTEUR/MANAGER).
    """
    serializer_class = StepCommentSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return StepComment.objects.filter(step_id=self.kwargs['pk']).select_related('author')

    def create(self, request, *args, **kwargs):
        step = get_object_or_404(ProjectStep, pk=self.kwargs['pk'])
        user = request.user
        can = (
            user.is_superuser
            or user.role in ('DIRECTEUR', 'MANAGER')
            or user.role in (step.responsables_roles or [])
        )
        if not can:
            return Response(
                {'detail': 'Vous ne pouvez pas commenter cette étape.'},
                status=status.HTTP_403_FORBIDDEN,
            )
        return super().create(request, *args, **kwargs)

    def perform_create(self, serializer):
        step = get_object_or_404(ProjectStep, pk=self.kwargs['pk'])
        serializer.save(step=step, author=self.request.user)


# ---------------------------------------------------------------------------
# API DRF — Commentaires projet
# ---------------------------------------------------------------------------

class ProjectCommentListCreateView(generics.ListCreateAPIView):
    """GET/POST /api/projets/<pk>/commentaires/"""
    serializer_class = ProjectCommentSerializer
    permission_classes = [IsAuthenticated, IsProjectMember]

    def get_queryset(self):
        return ProjectComment.objects.filter(project_id=self.kwargs['pk']).select_related('author')

    def perform_create(self, serializer):
        project = get_object_or_404(Project, pk=self.kwargs['pk'])
        serializer.save(project=project, author=self.request.user)
